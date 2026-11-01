"""结算单 Excel（发给代理）和学生名单导入（兼容现在的 Student Roster 格式）。"""
import io
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from commission import PROGRAMS


def statement_xlsx(agent, rec, lines, total):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\\/*?:\[\]]", "-", rec["term"])[:31]
    bold = Font(bold=True)
    ws.append(["Johnathan Academy – Agent Commission Statement"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Agent", f"{agent.get('name', '')} ({agent.get('code', '')})"])
    ws.append(["Term", rec["term"]])
    ws.append([])
    head = ["Student Name", "Program", "Courses", "Local / International", "Tuition Paid", "Rate", "Commission", "Note"]
    ws.append(head)
    for c in ws[5]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="DCE6F1")
    by_name = iter(rec.get("students", []))
    for l in lines:
        st = next(by_name, {}) if l.get("student") is not None else {}
        ws.append([l.get("student") or "", PROGRAMS.get(l.get("kind"), l["label"]), st.get("courses", ""),
                   st.get("origin", ""), l.get("fee"), f"{l['rate']:g}%" if l.get("rate") is not None else "",
                   l["amount"], st.get("note", "")])
    ws.append([])
    ws.append(["", "", "", "", "", "Total", total])
    ws.cell(ws.max_row, 6).font = bold
    ws.cell(ws.max_row, 7).font = bold
    for row in ws.iter_rows(min_row=6):
        for c in (row[4], row[6]):
            c.number_format = '"$"#,##0.00'
    for col, w in zip("ABCDEFGH", (26, 20, 34, 18, 14, 8, 14, 30)):
        ws.column_dimensions[col].width = w
    ws["A1"].alignment = Alignment(horizontal="left")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _program(courses):
    t = courses.lower()
    if "full-time" in t or "full time" in t:
        if "second" in t:
            return "FT2"
        if "third" in t:
            return "FT3"
        return "FT1"
    if re.search(r"\bap\b", t):
        return "AP"
    return "PT"


def parse_roster(data, sheet=None):
    """读取现在的 Student Roster Excel（每个学期一个工作表），返回 {工作表名: [学生行]}。"""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out = {}
    for ws in wb.worksheets:
        if sheet and ws.title != sheet:
            continue
        rows = list(ws.iter_rows(values_only=True))
        if not rows or not rows[0] or str(rows[0][0]).strip() != "Student Name":
            continue
        head = [str(h or "").strip() for h in rows[0]]
        idx = {h: i for i, h in enumerate(head)}
        students = []
        for r in rows[1:]:
            name = str(r[0] or "").strip()
            if not name or name.lower().startswith("total"):
                continue
            courses = ", ".join(str(r[idx[c]]).strip() for c in ("Course 1", "Course 2", "Course 3")
                                if c in idx and r[idx[c]] not in (None, ""))
            tuition = r[idx["Tuition Payment Total"]] if "Tuition Payment Total" in idx else None
            if not isinstance(tuition, (int, float)):
                continue
            status = str(r[idx["Student Status"]] or "").strip() if "Student Status" in idx else ""
            students.append({"name": name, "program": _program(courses), "courses": courses, "tuition": float(tuition),
                             "origin": status, "period": str(r[idx["Marks period"]] or "") if "Marks period" in idx else "",
                             "note": str(r[idx["Note"]] or "") if "Note" in idx else ""})
        out[ws.title] = students
    return out


# ---------- 发给代理的 Student Roster：以前所有学期 + 本学期新加一个工作表 ----------
ROSTER_HEAD = ["Student Name", "Course 1", "Course 2", "Course 3", "Tuition Payment Total",
               "Commission                                                     {rule}", "Agent", "Student Status",
               "Marks period", "Note"]
ORDINAL = {"FT1": "First", "FT2": "Second", "FT3": "Third"}
MONEY_FMT = '[$$-1009]#,##0.00'


def sheet_title(term):
    return re.sub(r"[\\/*?:\[\]]", "-", term)[:31]


def roster_rule(scheme):
    """表头里的比例说明，和现在 Roster 的写法一样，例如 Full 20%, 10%, 5% Part 15%。"""
    y = [scheme.get("y1"), scheme.get("y2"), scheme.get("y3")]
    full = f"{y[0]:g}%" if y[0] == y[1] == y[2] else ", ".join(f"{v:g}%" for v in y)
    if scheme.get("y1_6") and scheme["y1_6"] != scheme["y1"]:
        full += f" (6+ students {scheme['y1_6']:g}%)"
    return f"Full {full} Part {scheme.get('pt', 15):g}%"


def _courses(st):
    if st.get("program") in ORDINAL:
        g = str(st.get("grade") or "").strip().upper().lstrip("G")
        return [f"{'G' + g + ' ' if g else ''}Full-time {ORDINAL[st['program']]} year", "", ""]
    parts = [p.strip() for p in (st.get("courses") or "").split(",") if p.strip()]
    if st.get("program") == "AP" and not parts:
        parts = ["AP"]
    parts = parts[:3] if len(parts) <= 3 else parts[:2] + [", ".join(parts[2:])]
    return parts + [""] * (3 - len(parts))


def roster_rows(rec, lines):
    """结算单 → Roster 行：[(学生名, 课程1-3, 学费, 比例%, 佣金, 身份, 期间, 备注)]"""
    by_student = {}
    for l in lines:
        if l.get("student") is not None:
            by_student.setdefault(l["student"], []).append(l)
    rows = []
    for st in rec.get("students", []):
        l = (by_student.get(st.get("name")) or [None]).pop(0)
        if not l:
            continue
        rows.append({"name": st["name"], "courses": _courses(st), "fee": l["fee"], "rate": l["rate"], "amount": l["amount"],
                     "origin": st.get("origin", ""), "period": st.get("period", ""), "note": st.get("note", "")})
    for l in lines:
        if l.get("student") is None and l.get("fee") is None:  # 满 6 人补差、其它调整
            rows.append({"name": "", "courses": ["", "", ""], "fee": None, "rate": None, "amount": l["amount"],
                         "origin": "", "period": "", "note": l["label"]})
    return rows


def _blank_roster(title):
    from openpyxl.styles import Font as F
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(ROSTER_HEAD)
    for c in ws[1]:
        c.font = F(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 31.25
    for col, w in zip("ABCDEFGHIJ", (22, 26, 16, 16, 22, 32, 15, 15, 30, 18)):
        ws.column_dimensions[col].width = w
    return wb, ws


def roster_workbook(base, term, rec, lines, scheme, agent_label, today=None):
    """base：代理以前的 Roster（字节，可为 None）。返回加上（或替换）本学期工作表后的 Excel 字节。
    新工作表照最后一个学期的样式复制，所以字体、列宽、金额格式都和以前一样。"""
    import datetime as _dt
    today = today or _dt.date.today()
    title = sheet_title(term)
    wb = openpyxl.load_workbook(io.BytesIO(base)) if base else None
    if wb:
        tpl = next((ws for ws in reversed(wb.worksheets)
                    if ws.title != title and str(ws["A1"].value or "").strip() == "Student Name"), None)
        if title in wb.sheetnames:
            idx = wb.sheetnames.index(title)
            del wb[title]
        else:
            idx = len(wb.sheetnames)
        if tpl is not None:
            ws = wb.copy_worksheet(tpl)
            ws.title = title
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.value = None
        else:
            ws = wb.create_sheet(title)
            ws.append(ROSTER_HEAD)
        wb.move_sheet(ws, offset=idx - wb.sheetnames.index(title))
        for i, h in enumerate(ROSTER_HEAD, 1):  # 以前有的学期少一栏或多一栏，新表一律用标准的 10 栏
            ws.cell(1, i).value = h
        for i in range(len(ROSTER_HEAD) + 1, ws.max_column + 1):
            ws.cell(1, i).value = None
        wb.active = wb.sheetnames.index(title)
        for s in wb.worksheets:
            s.sheet_view.tabSelected = s.title == title
    else:
        wb, ws = _blank_roster(title)
    ws.cell(1, 6).value = ROSTER_HEAD[5].format(rule=roster_rule(scheme))
    rows = roster_rows(rec, lines)
    r = 1
    for r, x in enumerate(rows, 2):
        vals = [x["name"], *x["courses"], x["fee"],
                f"=E{r}*{x['rate'] / 100:g}" if x["rate"] is not None and x["fee"] is not None else x["amount"],
                agent_label if x["name"] else "", x["origin"], x["period"], x["note"]]
        for i, v in enumerate(vals, 1):
            ws.cell(r, i).value = v if v != "" else None
        ws.cell(r, 5).number_format = MONEY_FMT
        ws.cell(r, 6).number_format = MONEY_FMT
    total_row = max(13, r + 3)
    ws.cell(total_row, 5).value = f"Total Commission Until {today.strftime('%b')}-{today.day}-{today.year}"
    ws.cell(total_row, 6).value = f"=SUM(F2:F{max(r, 2)})"
    ws.cell(total_row, 5).font = Font(bold=True)
    ws.cell(total_row, 6).font = Font(bold=True)
    ws.cell(total_row, 6).number_format = MONEY_FMT
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def last_agent_label(base):
    """以前 Roster 里 Agent 一栏最后写的称呼（例如 Shanice、Jerry），新工作表沿用。"""
    if not base:
        return ""
    wb = openpyxl.load_workbook(io.BytesIO(base), data_only=True)
    for ws in reversed(wb.worksheets):
        head = [str(c.value or "").strip() for c in ws[1]]
        if "Agent" not in head:
            continue
        col = head.index("Agent")
        vals = [r[col] for r in ws.iter_rows(min_row=2, values_only=True) if len(r) > col and r[col]]
        if vals:
            return str(vals[-1]).strip()
    return ""


_CELL = re.compile(r"^([A-Z]+)(\d+)$")


def _eval(formula, grid):
    """预览用：只算 Roster 里常见的 =E2*0.25、=SUM(F2:F9)、=E2+E3 这几种；算不了就显示公式。"""
    from openpyxl.utils import column_index_from_string as ci

    def val(ref):
        m = _CELL.match(ref)
        v = grid.get((int(m.group(2)), ci(m.group(1)))) if m else None
        return v if isinstance(v, (int, float)) else 0
    f = formula[1:].replace(" ", "")
    m = re.fullmatch(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", f, re.I)
    if m:
        c1, r1, c2, r2 = ci(m.group(1).upper()), int(m.group(2)), ci(m.group(3).upper()), int(m.group(4))
        return sum(v for (r, c), v in grid.items() if r1 <= r <= r2 and c1 <= c <= c2 and isinstance(v, (int, float)))
    if re.fullmatch(r"[A-Z]+\d+([*+\-/][A-Z0-9.]+)*", f):
        expr = re.sub(r"[A-Z]+\d+", lambda x: repr(val(x.group(0))), f)
        try:
            return eval(expr, {"__builtins__": {}}, {})  # 只剩数字和 + - * /
        except Exception:  # noqa: BLE001
            return formula
    return formula


def preview(data, max_rows=60):
    """把 Excel 变成网页预览：[{title, rows: [[文字…]]}]，公式按上面的规则算出来。"""
    wb = openpyxl.load_workbook(io.BytesIO(data))
    cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out = []
    for ws in wb.worksheets:
        cw = cached[ws.title]
        grid = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    cv = cw.cell(c.row, c.column).value
                    v = cv if cv is not None else None
                grid[(c.row, c.column)] = v
        for row in ws.iter_rows():  # 第二遍：没有缓存值的公式自己算（SUM 用到的先算出来）
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and grid[(c.row, c.column)] is None:
                    grid[(c.row, c.column)] = _eval(c.value, grid)
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=SUM") and cw.cell(c.row, c.column).value is None:
                    grid[(c.row, c.column)] = _eval(c.value, grid)
        n_rows = max([r for (r, c), v in grid.items() if v not in (None, "")] or [0])
        n_cols = max([c for (r, c), v in grid.items() if v not in (None, "")] or [0])
        rows = []
        for r in range(1, min(n_rows, max_rows) + 1):
            cells = []
            for c in range(1, n_cols + 1):
                v = grid.get((r, c))
                fmt = ws.cell(r, c).number_format or ""
                if isinstance(v, (int, float)) and ("$" in fmt or (r > 1 and c in (5, 6))):
                    v = f"${v:,.2f}"
                elif isinstance(v, float):
                    v = f"{v:g}"
                cells.append("" if v is None else str(v))
            rows.append(cells)
        out.append({"title": ws.title, "rows": rows, "more": n_rows > max_rows})
    return out
