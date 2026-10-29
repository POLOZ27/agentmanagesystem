"""年度付款汇总和 T4A。

按「付款日期」所在的自然年统计（T4A 按日历年报，不按学年）。
T4A 规则（给会计核对用，最终以会计意见为准）：
  - 只有在加拿大的代理才需要：个人 / 独资（未注册公司）全年付款超过 $500 → 开 T4A，
    佣金填 Box 020（Self-employed commissions）。
  - 注册公司（Inc. / Ltd. / Corp.）一般不需要 T4A。
  - 加拿大以外的代理，服务不在加拿大境内提供，不需要 T4A。
  - 截止日期：次年 2 月最后一天前发给代理并报给 CRA。
"""
import calendar
import datetime as dt
import io
import re

import openpyxl
from openpyxl.styles import Font, PatternFill

THRESHOLD = 500
ENTITY_TYPES = {"individual": "个人 / 独资", "corporation": "注册公司"}
CORP_RE = re.compile(r"\b(inc|ltd|corp|corporation|limited|incorporated|ltée)\b\.?", re.I)


def is_canadian(a):
    return (a.get("country") or "").strip().lower() in ("canada", "加拿大", "ca")


def entity_type(a):
    if a.get("entity_type") in ENTITY_TYPES:
        return a["entity_type"]
    name = " ".join([a.get("legal_name") or "", a.get("payee") or "", a.get("name") or ""])
    return "corporation" if CORP_RE.search(name) else "individual"


def deadline(year):
    return dt.date(year + 1, 2, calendar.monthrange(year + 1, 2)[1])


def summarize(year, rows, agents):
    """rows: enrich() 过的结算单。返回按代理汇总的列表（只含当年有付款的代理）。"""
    by = {}
    for r in rows:
        if r.get("stage") != "已付" or not (r.get("paid_on") or "").startswith(str(year)):
            continue
        a = agents.get(r["agent_id"], {})
        s = by.setdefault(r["agent_id"], {"agent": a, "payments": [], "total": 0.0})
        s["payments"].append({"date": r["paid_on"], "term": r["term"], "amount": r["amount"],
                              "method": r.get("pay_method") or "", "ref": r.get("pay_ref") or ""})
        s["total"] = round(s["total"] + r["amount"], 2)
    out = []
    for s in by.values():
        a = s["agent"]
        s["payments"].sort(key=lambda p: p["date"])
        s["canadian"] = is_canadian(a)
        s["entity"] = entity_type(a)
        s["t4a"] = s["canadian"] and s["entity"] == "individual" and s["total"] > THRESHOLD
        if not s["canadian"]:
            s["t4a_note"] = "加拿大以外，不需要"
        elif s["entity"] == "corporation":
            s["t4a_note"] = "注册公司，一般不需要（请会计确认）"
        elif s["total"] <= THRESHOLD:
            s["t4a_note"] = f"全年未超过 ${THRESHOLD}，不需要"
        else:
            s["t4a_note"] = "需要开 T4A（Box 020）" + ("" if a.get("tax_number") else "，缺 SIN / BN")
        out.append(s)
    return sorted(out, key=lambda s: (not s["t4a"], -s["total"]))


def years(rows):
    ys = {int(r["paid_on"][:4]) for r in rows if r.get("stage") == "已付" and r.get("paid_on")}
    ys.add(dt.date.today().year)
    return sorted(ys, reverse=True)


def workbook(year, summary, payer):
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DCE6F1")

    def header(ws, cols):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font = bold
            c.fill = head_fill

    ws = wb.active
    ws.title = f"{year} 付款汇总"
    header(ws, ["代理编号", "代理名称", "收款人 / 支票抬头", "国家", "类型", "付款次数", "全年合计", "T4A", "付款明细"])
    for s in summary:
        a = s["agent"]
        ws.append([a.get("code"), a.get("name"), a.get("payee") or a.get("legal_name") or a.get("name"),
                   a.get("country"), ENTITY_TYPES[s["entity"]], len(s["payments"]), s["total"], s["t4a_note"],
                   "; ".join(f"{p['date']} {p['term']} ${p['amount']:,.2f} {p['method']}" for p in s["payments"])])
    ws.append([])
    ws.append(["", "", "", "", "", "合计", round(sum(s["total"] for s in summary), 2)])
    for col, w in zip("ABCDEFGHI", (11, 30, 30, 12, 12, 9, 13, 34, 80)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[6].number_format = '"$"#,##0.00'

    t = wb.create_sheet("T4A")
    t.append([f"T4A {year} – Payer: {payer.get('name', '')}  BN: {payer.get('bn', '')}  "
              f"Deadline: {deadline(year).isoformat()}"])
    t["A1"].font = Font(bold=True, size=12)
    header(t, ["Recipient name (as on slip)", "SIN / BN", "Address", "City / Province / Postal code",
               "Box 020 Self-employed commissions", "Box 048 Fees for services", "Agent No.", "Note"])
    for s in summary:
        if not s["t4a"]:
            continue
        a = s["agent"]
        t.append([a.get("payee") or a.get("legal_name") or a.get("name"), a.get("tax_number") or "",
                  a.get("address") or "", a.get("city") or "", s["total"], None, a.get("code"),
                  "" if a.get("tax_number") else "Missing SIN/BN – request from agent"])
    for col, w in zip("ABCDEFGH", (32, 16, 36, 28, 18, 14, 11, 34)):
        t.column_dimensions[col].width = w
    for row in t.iter_rows(min_row=3):
        row[4].number_format = '"$"#,##0.00'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def agent_statement(year, s):
    """发给单个代理的年度付款明细（Excel）。"""
    a = s["agent"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{year}"
    ws.append([f"Johnathan Academy – {year} Annual Commission Payment Summary"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Agent", f"{a.get('name', '')} ({a.get('code', '')})"])
    ws.append(["Payee", a.get("payee") or a.get("legal_name") or a.get("name")])
    ws.append([])
    ws.append(["Payment date", "Term", "Method", "Reference", "Amount (CAD)"])
    for c in ws[5]:
        c.font = Font(bold=True)
    for p in s["payments"]:
        ws.append([p["date"], p["term"], p["method"], p["ref"], p["amount"]])
    ws.append([])
    ws.append(["", "", "", "Total", s["total"]])
    ws.cell(ws.max_row, 4).font = Font(bold=True)
    for row in ws.iter_rows(min_row=6):
        row[4].number_format = '"$"#,##0.00'
    for col, w in zip("ABCDE", (14, 26, 14, 18, 14)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
