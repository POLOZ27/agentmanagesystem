import datetime as dt
import io
import json
import os
import re
import secrets
import tempfile
import time
from functools import wraps

import openpyxl
import pyotp
from flask import (Flask, abort, flash, redirect, render_template, request, send_file, session, url_for)

import threading

import admissions
import ap_courses
import ai
import portal
import backup
import tax
import local_ai
import contracts
import email_templates as et
import mailer
import importer
import inbox
import outreach_i18n as i18n
import statements
from commission import (LEVELS, PROGRAMS, academic_year, agent_scheme, calculate, parse_scheme, scheme_label,
                        total_fees, y1_rate_used)
from vault import Vault

DATA_DIR = os.environ.get("AGENT_PORTAL_DATA", os.path.join(os.path.dirname(__file__), "data"))
IDLE_MINUTES = 30
STATUSES = ["潜在", "洽谈中", "已签约", "活跃", "不活跃", "终止"]
# 新代理从第一次联系到正式开通的 7 步
ONBOARDING = [
    ("contact", "初次联系", "发出开发邮件（在「开发新代理」里一键发送）"),
    ("meeting", "回复 / 线上会议", "代理回信或约了 30 分钟会议"),
    ("vetting", "申请表 + 背景核实", "收到代理申请表，完成 Agencies Reference Check"),
    ("level", "确定等级", "标准级 20/25% 或 高级 25/30%（第一年，满 6 人起高档）"),
    ("contract", "生成并发送合同", "系统按等级生成合同，连同 Start Letter 一起发出"),
    ("signed", "收回签字合同", "上传代理签回的合同和 Schedule A 银行信息"),
    ("activate", "开通", "确认编号，发欢迎信，加入 News 推送名单"),
]
STAGES = ["草稿", "已发结算单", "代理已确认", "已交会计", "已付"]
STUDENT_STAGES = ["收到申请", "资料齐全", "已发录取信", "已缴学费", "学签获批", "已入学", "未入学"]
STUDENT_HINTS = {
    "收到申请": "代理发来申请表", "资料齐全": "护照、成绩单等材料收齐", "已发录取信": "发出 Letter of Acceptance",
    "已缴学费": "学费到账", "学签获批": "国际生拿到学签（本地生可跳过）", "已入学": "开始上课，可以计入佣金结算单",
    "未入学": "放弃、拒签或转校",
}
PROSPECT_STAGES = ["未联系", "已发邮件", "已跟进", "已回复", "已约会议", "已转为代理", "不合适"]

app = Flask(__name__)
app.secret_key = secrets.token_bytes(32)  # 每次重启更换：重启后所有登录失效
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Strict",
                  SESSION_COOKIE_SECURE=os.environ.get("AGENT_PORTAL_HTTPS") == "1",
                  MAX_CONTENT_LENGTH=400 * 1024 * 1024)  # 首次导入要一次上传全部已签合同
vault = Vault(DATA_DIR)
_failed = []  # 登录失败时间戳，用于限流


# ---------- 登录 ----------
def login_required(f):
    @wraps(f)
    def wrapper(*a, **kw):
        if not vault.initialized:
            return redirect(url_for("setup"))
        last = session.get("seen", 0)
        if not session.get("ok") or not vault.unlocked or time.time() - last > IDLE_MINUTES * 60:
            session.clear()
            return redirect(url_for("login"))
        session["seen"] = time.time()
        return f(*a, **kw)
    return wrapper


@app.after_request
def headers(resp):
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["Cache-Control"] = "no-store"
    resp.headers["Referrer-Policy"] = "no-referrer"
    return resp


@app.before_request
def csrf():
    if request.method == "POST":
        if request.endpoint not in ("login", "setup") and request.form.get("_csrf") != session.get("csrf"):
            abort(400)


@app.context_processor
def ctx():
    session.setdefault("csrf", secrets.token_hex(16))
    n_inbox = sum(b["complete"] for b in inbox.bundles(vault)) if vault.unlocked and session.get("ok") else 0
    return {"csrf": session["csrf"], "scheme_label": scheme_label, "unlocked": vault.unlocked, "portal_on": portal.enabled(),
            "inbox_n": n_inbox, "parent_url": parent_url() if session.get("ok") else None,
            "mail_accounts": lambda: mailer.accounts(settings()), "MAIL_KINDS": mailer.MAIL_KINDS,
            "default_sender": lambda kind: (mailer.account_for(settings(), kind) or {}).get("user")}


# 顶部导航里的一级页面不显示「返回」；其余页面的返回默认回到上一页，没有上一页时回到这里
TOP_PAGES = {"dashboard", "agents_list", "students", "prospects", "commissions", "mail", "materials", "performance",
             "payments", "settings_page", "login", "setup", "recovery", "static"}
PARENT = {"commission": "commissions", "agent": "agents_list", "contract": "agents_list", "renew": "renewals",
          "renewals": "agents_list", "prospect": "prospects", "outreach": "prospects", "student": "students",
          "mail": "mail", "compose": "mail", "inbox": "dashboard", "trash": "settings_page"}


def parent_url():
    ep = request.endpoint or ""
    if ep in TOP_PAGES:
        return None
    args = request.view_args or {}
    if ep.startswith("commission") and args.get("cid"):
        c = vault.get("commission", args["cid"])
        return url_for("commissions", term=c["term"]) if c else url_for("commissions")
    if ep.startswith(("contract", "agent_")) and args.get("aid") and ep != "agent_detail":
        return url_for("agent_detail", aid=args["aid"])
    for k, v in PARENT.items():
        if ep.startswith(k):
            return url_for(v)
    return url_for("dashboard")


@app.route("/setup", methods=["GET", "POST"])
def setup():
    if vault.initialized:
        return redirect(url_for("login"))
    token = app.config.get("SETUP_TOKEN")  # 生产环境：只有拿到服务器日志里链接的人能做首次设置
    if token and request.values.get("t") != token:
        return "首次设置需要使用部署时显示的完整链接。", 403
    if request.method == "POST":
        pw = request.form["password"]
        if len(pw) < 12 or pw != request.form["password2"]:
            flash("主密码至少 12 位，且两次输入一致")
            return redirect(url_for("setup", t=token))
        secret = request.form["totp_secret"]
        if not pyotp.TOTP(secret).verify(request.form["code"], valid_window=1):
            flash("验证码不对，请在手机验证器里重新查看")
            return redirect(url_for("setup", t=token))
        recovery = vault.initialize(pw, secret)
        vault.unlock(vault.try_unwrap(pw))
        inbox.ensure_key(vault)
        session.update(ok=True, seen=time.time())
        return render_template("recovery.html", recovery=recovery)
    secret = pyotp.random_base32()
    uri = pyotp.TOTP(secret).provisioning_uri("owner", issuer_name="Agent Portal")
    return render_template("setup.html", secret=secret, uri=uri)


# ---------- 信任的设备：密码还是每次都要，只是 30 天内不再要验证码 ----------
TRUST_DAYS = 30
TRUST_COOKIE = "ap_trusted"
TRUST_FILE = os.path.join(DATA_DIR, "trusted_devices.json")  # 只存令牌的哈希，没有任何代理数据


def _trusted_load():
    try:
        with open(TRUST_FILE) as fh:
            rows = json.load(fh)
    except (OSError, ValueError):
        return {}
    now = time.time()
    return {k: v for k, v in rows.items() if v.get("exp", 0) > now}


def _trusted_save(rows):
    tmp = TRUST_FILE + ".tmp"
    with open(tmp, "w") as fh:
        json.dump(rows, fh)
    os.chmod(tmp, 0o600)
    os.replace(tmp, TRUST_FILE)


def _token_hash(tok):
    import hashlib
    return hashlib.sha256(tok.encode()).hexdigest()


def device_trusted():
    tok = request.cookies.get(TRUST_COOKIE)
    return bool(tok) and _token_hash(tok) in _trusted_load()


@app.route("/login", methods=["GET", "POST"])
def login():
    if not vault.initialized:
        return redirect(url_for("setup"))
    if request.method == "POST":
        now = time.time()
        _failed[:] = [t for t in _failed if now - t < 900]
        if len(_failed) >= 5:
            flash("失败次数过多，请 15 分钟后再试")
            return redirect(url_for("login"))
        dek = vault.try_unwrap(request.form["password"])
        trusted = device_trusted()
        if not dek or not (trusted or pyotp.TOTP(vault.totp_secret(dek)).verify(request.form.get("code", ""), valid_window=1)):
            _failed.append(now)
            vault.audit("login failed")
            flash("密码或验证码错误")
            return redirect(url_for("login"))
        if not vault.unlocked:
            vault.unlock(dek)
        inbox.ensure_key(vault)
        vault.audit("login ok")
        try:
            refresh_old_drafts()
        except Exception:  # 清理失败不能挡住登录
            app.logger.exception("refresh_old_drafts")
        session.clear()
        session.update(ok=True, seen=now, csrf=secrets.token_hex(16))
        resp = redirect(url_for("dashboard"))
        if not trusted and request.form.get("remember"):
            tok = secrets.token_urlsafe(32)
            rows = _trusted_load()
            rows[_token_hash(tok)] = {"exp": now + TRUST_DAYS * 86400, "added": dt.date.today().isoformat(),
                                      "ua": (request.user_agent.string or "")[:120]}
            _trusted_save(rows)
            https = request.is_secure or request.headers.get("X-Forwarded-Proto") == "https"
            resp.set_cookie(TRUST_COOKIE, tok, max_age=TRUST_DAYS * 86400, httponly=True, secure=https, samesite="Strict")
            vault.audit("device trusted for 30 days")
        return resp
    return render_template("login.html", trusted=device_trusted(), trust_days=TRUST_DAYS)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.post("/lock")
@login_required
def lock():
    vault.lock()
    session.clear()
    flash("保险库已锁定：内存中的密钥已清除，定时邮件暂停，需要重新登录解锁。")
    return redirect(url_for("login"))


# ---------- 工具 ----------
def agents_by_id():
    return {a["id"]: a for a in vault.all("agent")}


def terms():
    """所有学期，按先后排：结算单里出现过的 + 自己新建的（还没有结算单的学期）。"""
    seen = {}
    for c in vault.all("commission"):
        seen[c["term"]] = max(seen.get(c["term"], -1), c.get("term_order", 999))
    for t in vault.all("term"):
        seen.setdefault(t["name"], t.get("order", 999))
    return [t for t, _ in sorted(seen.items(), key=lambda x: x[1])]


def next_terms(existing):
    """新建学期时的建议名称：本学年和下一学年里还没有的学期，格式和佣金总表的工作表名一致。"""
    today = dt.date.today()
    y = today.year % 100 if today.month >= 9 else today.year % 100 - 1
    out = []
    for a in (y, y + 1):
        ay = f"{a:02d}-{a + 1:02d}"
        out += [f"Sep-Dec Fall {ay}", f"Jan-Apr Winter {ay}", f"May-June Spring {ay}", f"20{a + 1:02d} Summer School"]
    return [t for t in out if t not in existing]


def stage_of(c):
    s = c.get("status") or "草稿"
    return {"待付": "草稿"}.get(s, s)


def ay_group(c, all_comms=None):
    ay = academic_year(c["term"])
    pool = all_comms if all_comms is not None else vault.all("commission")
    return [x for x in pool if x["agent_id"] == c["agent_id"] and academic_year(x["term"]) == ay]


def enrich(c, agents, all_comms=None):
    a = agents.get(c["agent_id"], {})
    scheme = agent_scheme(a) if not c.get("rates") else agent_scheme({"rates": c["rates"]})
    group = ay_group(c, all_comms)
    lines, total = calculate(c, scheme, group)
    c = dict(c, agent=a, lines=lines, calc=total, fees=total_fees(c), scheme=scheme, stage=stage_of(c))
    ex = c.get("excel_commission")
    c["mismatch"] = c.get("imported") and ex is not None and abs(ex - total) >= 1
    c["amount"] = ex if c.get("imported") and ex is not None else total
    c["ft1_ay"] = len({(s.get("name") or "").lower() for x in group for s in x.get("students", []) if s.get("program") == "FT1"})
    return c


def days_left(d):
    if not d:
        return None
    return (dt.date.fromisoformat(d) - dt.date.today()).days


def next_code():
    year = dt.date.today().year
    nums = [int(a["code"][5:]) for a in vault.all("agent")
            if a.get("code") and a["code"].startswith(f"J{year}") and a["code"][5:].isdigit()]
    return f"J{year}{(max(nums) + 1 if nums else 1):02d}"


def fnum(v):
    v = (v or "").replace(",", "").replace("$", "").strip()
    return float(v) if v else None


def money(v):
    if not isinstance(v, (int, float)):
        return ""
    return f"${v:,.0f}" if float(v).is_integer() else f"${v:,.2f}"


def onboarding_step(a):
    done = a.get("onboarding") or {}
    for i, (key, _, _) in enumerate(ONBOARDING):
        if key not in done:
            return i
    return len(ONBOARDING)


def mark_step(a, key):
    a.setdefault("onboarding", {})
    a["onboarding"].setdefault(key, dt.date.today().isoformat())


app.jinja_env.filters["money"] = money
app.jinja_env.globals.update(days_left=days_left, agent_scheme=agent_scheme, STATUSES=STATUSES, STAGES=STAGES, LEVELS=LEVELS, PROGRAMS=PROGRAMS,
                             ONBOARDING=ONBOARDING, onboarding_step=onboarding_step, PROSPECT_STAGES=PROSPECT_STAGES,
                             COUNTRY_ZH=i18n.COUNTRY_ZH, STUDENT_STAGES=STUDENT_STAGES, STUDENT_HINTS=STUDENT_HINTS,
                             LANG_LABELS={k: v["label"] for k, v in i18n.TRANSLATIONS.items()}, lang_for=i18n.lang_for)


def settings():
    return mailer.settings(vault)


def save_upload(field):
    f = request.files.get(field)
    if not f or not f.filename:
        return None
    return {"name": f.filename, "fid": vault.put_file(f.read()), "uploaded": dt.date.today().isoformat()}


@app.route("/files/<fid>")
@login_required
def file_download(fid):
    if not re.fullmatch(r"[0-9a-f]{32}", fid):
        abort(404)
    vault.audit("file downloaded")
    return send_file(io.BytesIO(vault.get_file(fid)), download_name=request.args.get("name", "file"), as_attachment=True)


# ---------- 总览 ----------
@app.route("/")
@login_required
def dashboard():
    agents = agents_by_id()
    ts = terms()
    cur = ts[-1] if ts else None
    comms = vault.all("commission")
    rows = [enrich(c, agents, comms) for c in comms]
    cur_rows = [r for r in rows if r["term"] == cur]
    ay = current_ay()
    ay_rows = [r for r in rows if academic_year(r["term"]) == ay]
    open_rows = [r for r in rows if r["stage"] in STAGES[:-1] and not r.get("imported")]
    expiring = sorted([a for a in agents.values() if a.get("contract_expiry") and a.get("status") in ("活跃", "已签约")
                       and days_left(a["contract_expiry"]) <= 120], key=lambda a: a["contract_expiry"])
    pipeline = sorted([a for a in agents.values() if a.get("status") in ("潜在", "洽谈中", "已签约")],
                      key=lambda a: -onboarding_step(a))
    by_term = [(t, sum(r["amount"] for r in rows if r["term"] == t)) for t in ts[-8:]]
    prospects = vault.all("prospect")
    studs = vault.all("student")
    t4a_due = None
    today = dt.date.today()
    if today.month <= 2:  # 1–2 月提醒上一年的 T4A
        need = [x for x in tax.summarize(today.year - 1, rows, agents) if x["t4a"]]
        if need:
            t4a_due = {"year": today.year - 1, "n": len(need), "deadline": tax.deadline(today.year - 1)}
    return render_template("dashboard.html", agents=agents, cur=cur, cur_rows=cur_rows, t4a_due=t4a_due,
                           open_students=sum(1 for x in studs if x.get("stage") not in ("已入学", "未入学")),
                           enrolled_now=sum(1 for x in studs if x.get("stage") == "已入学" and academic_year(x.get("intake") or "") == ay),
                           ay=ay, ay_total=sum(r["amount"] for r in ay_rows), ay_agents=len({r["agent_id"] for r in ay_rows}),
                           cur_total=sum(r["amount"] for r in cur_rows), open_rows=open_rows,
                           stage_counts={s: [r for r in open_rows if r["stage"] == s] for s in STAGES[:-1]},
                           expiring=expiring, pipeline=pipeline, by_term=by_term,
                           active=sum(1 for a in agents.values() if a.get("status") == "活跃"),
                           prospects=prospects,
                           followups=[p for p in prospects if needs_followup(p)],
                           inbox=sorted([m for m in vault.all("mail") if m["direction"] == "in"],
                                        key=lambda m: m["at"], reverse=True)[:5],
                           campaigns=[c for c in vault.all("campaign") if c.get("active")])


# ---------- 代理 ----------
@app.route("/agents")
@login_required
def agents_list():
    q = request.args.get("q", "").strip().lower()
    st = request.args.get("status", "")
    by_id = agents_by_id()
    comms = vault.all("commission")
    earned = {}
    for c in comms:
        earned[c["agent_id"]] = earned.get(c["agent_id"], 0) + enrich(c, by_id, comms)["amount"]
    agents = sorted(vault.all("agent"), key=lambda a: a.get("code") or "J9" + a["name"])
    if q:
        agents = [a for a in agents if q in " ".join(str(a.get(k, "")) for k in
                                                     ("name", "code", "email", "contact", "country", "payee")).lower()]
    if st:
        agents = [a for a in agents if a.get("status") == st]
    return render_template("agents.html", agents=agents, earned=earned, q=q, st=st)


def agent_from_form(a, f):
    for k in ("name", "legal_name", "code", "contact", "email", "phone", "country", "city", "address", "website",
              "payee", "etransfer", "detail", "notes", "status", "contract_expiry", "source", "level",
              "entity_type", "tax_number"):
        if k in f:
            a[k] = f.get(k, "").strip()
    a["contract_expiry"] = a.get("contract_expiry") or None
    if a.get("level") in LEVELS:
        a["rates"] = [LEVELS[a["level"]]["y1"], LEVELS[a["level"]]["y2"], LEVELS[a["level"]]["y3"]]
    else:
        a["level"] = ""
        rates = [f.get("rate_y1"), f.get("rate_y2"), f.get("rate_y3")]
        if all(rates):
            a["rates"] = [float(r) for r in rates]
    return a


@app.route("/agents/new", methods=["GET", "POST"])
@login_required
def agent_new():
    if request.method == "POST":
        a = agent_from_form({"created": dt.date.today().isoformat(), "contracts": [], "onboarding": {}}, request.form)
        a = vault.put("agent", a)
        flash(f"已新建代理 {a['name']}")
        return redirect(url_for("agent_detail", aid=a["id"]))
    return render_template("agent_form.html", a={"code": next_code(), "status": "洽谈中", "level": "standard"}, new=True)


@app.route("/agents/<aid>", methods=["GET", "POST"])
@login_required
def agent_detail(aid):
    a = vault.get("agent", aid) or abort(404)
    if request.method == "POST":
        agent_from_form(a, request.form)
        vault.put("agent", a)
        flash("已保存")
        return redirect(url_for("agent_detail", aid=aid))
    comms = vault.all("commission")
    mine = sorted([enrich(c, {aid: a}, comms) for c in comms if c["agent_id"] == aid],
                  key=lambda c: c.get("term_order", 999), reverse=True)
    mails = sorted([m for m in vault.all("mail") if m.get("agent_id") == aid], key=lambda m: m["at"], reverse=True)
    start, end = contracts.default_period()
    studs = sorted([x for x in vault.all("student") if x.get("agent_id") == aid], key=lambda x: x.get("intake") or "", reverse=True)
    return render_template("agent_detail.html", a=a, comms=mine, mails=mails, total=sum(c["amount"] for c in mine), students=studs,
                           step=onboarding_step(a), start=start, end=end)


# ---------- 回收站：删掉的合同和代理先保留 30 天，可以恢复，之后自动清除 ----------
TRASH_DAYS = 30


def _remove_file(fid):
    if fid:
        try:
            os.remove(os.path.join(DATA_DIR, "files", fid))
        except OSError:
            pass


def trash_contract(a, c, why):
    """把一份合同放进回收站（文件留在原地，不删）。"""
    if c.get("fid"):
        vault.put("trash", {"type": "contract", "agent_id": a["id"], "label": f"{a.get('name', '')} · {c.get('name', '')}",
                            "why": why, "c": c, "at": time.time()})


def purge_trash():
    """超过 30 天的回收站内容真正删除。"""
    cutoff = time.time() - TRASH_DAYS * 86400
    for t in vault.all("trash"):
        if t.get("at", 0) < cutoff:
            for fid in t.get("fids") or [t.get("c", {}).get("fid")]:
                _remove_file(fid)
            vault.delete("trash", t["id"])


@app.route("/trash")
@login_required
def trash_list():
    rows = sorted(vault.all("trash"), key=lambda t: -t.get("at", 0))
    for t in rows:
        t["when"] = dt.datetime.fromtimestamp(t.get("at", 0)).strftime("%Y-%m-%d %H:%M")
        t["left"] = max(0, TRASH_DAYS - int((time.time() - t.get("at", 0)) // 86400))
    return render_template("trash.html", rows=rows, days=TRASH_DAYS)


@app.post("/trash/<tid>/restore")
@login_required
def trash_restore(tid):
    t = vault.get("trash", tid) or abort(404)
    if t["type"] == "contract":
        a = vault.get("agent", t["agent_id"])
        if not a:
            flash("这个代理已经不在了，请先从回收站恢复代理")
            return redirect(url_for("trash_list"))
        c = dict(t["c"], kind="已恢复" if t["c"].get("kind") in ("续签待签", "待签（系统生成）") else t["c"].get("kind"))
        a.setdefault("contracts", []).append(c)
        vault.put("agent", a)
    else:
        vault.put("agent", t["agent"])
        for kind, recs in (t.get("records") or {}).items():
            for r in recs:
                vault.put(kind, r)
    vault.delete("trash", tid)
    vault.audit("restored from trash")
    flash(f"已恢复：{t.get('label', '')}")
    return redirect(url_for("trash_list"))


@app.post("/trash/<tid>/purge")
@login_required
def trash_purge(tid):
    t = vault.get("trash", tid) or abort(404)
    for fid in t.get("fids") or [t.get("c", {}).get("fid")]:
        _remove_file(fid)
    vault.delete("trash", tid)
    vault.audit("trash item purged")
    flash("已彻底删除")
    return redirect(url_for("trash_list"))


@app.post("/agents/<aid>/delete")
@login_required
def agent_delete(aid):
    """删除导入错了的代理：连同它的结算单、学生申请和合同文件一起删掉。"""
    a = vault.get("agent", aid) or abort(404)
    records = {}
    for kind in ("commission", "student"):
        for r in vault.all(kind):
            if r.get("agent_id") == aid:
                records.setdefault(kind, []).append(r)
                vault.delete(kind, r["id"])
    vault.put("trash", {"type": "agent", "label": a.get("name", ""), "why": "删除代理", "agent": a, "records": records,
                        "fids": [c["fid"] for c in a.get("contracts") or [] if c.get("fid")], "at": time.time()})
    vault.delete("agent", aid)
    vault.audit("agent deleted (to trash)")
    flash(f"已删除代理 {a.get('name', '')}，连同结算单、学生申请和合同一起放进回收站，{TRASH_DAYS} 天内可以恢复")
    return redirect(url_for("agents_list"))


@app.route("/agents/<aid>/edit")
@login_required
def agent_edit(aid):
    return render_template("agent_form.html", a=vault.get("agent", aid) or abort(404), new=False)


@app.post("/agents/<aid>/onboarding")
@login_required
def agent_onboarding(aid):
    a = vault.get("agent", aid) or abort(404)
    old = a.get("onboarding") or {}
    a["onboarding"] = {k: old.get(k) or dt.date.today().isoformat() for k, _, _ in ONBOARDING if request.form.get(k)}
    if "signed" in a["onboarding"] and a.get("status") in ("潜在", "洽谈中"):
        a["status"] = "已签约"
    vault.put("agent", a)
    return redirect(url_for("agent_detail", aid=aid))


@app.post("/agents/<aid>/contract")
@login_required
def contract_generate(aid):
    """第 5 步：按等级生成合同；可直接连同 Start Letter 发给代理。"""
    a = vault.get("agent", aid) or abort(404)
    level = request.form.get("level") or a.get("level") or "standard"
    start = dt.date.fromisoformat(request.form["start"])
    end = dt.date.fromisoformat(request.form["end"])
    if not a.get("code"):
        a["code"] = next_code()
    a["level"] = level
    a["rates"] = [LEVELS[level]["y1"], LEVELS[level]["y2"], LEVELS[level]["y3"]]
    fname, data = contracts.generate(a, level, start, end)
    a.setdefault("contracts", []).append({"name": fname, "fid": vault.put_file(data), "expiry": end.isoformat(),
                                          "uploaded": dt.date.today().isoformat(), "kind": "待签（系统生成）",
                                          "start": start.isoformat(), "tpl": TEMPLATE_REV})
    mark_step(a, "level")
    if request.form.get("send") and a.get("email"):
        tpl = et.fill(et.START_LETTER, agent=a.get("legal_name") or a["name"], contact=a.get("contact") or a["name"],
                      kit_link=kit_link())
        with mailer.Outbox(vault) as box:
            box.send(a["email"], tpl["subject"], tpl["body"], [(fname, data)], agent_id=aid, kind="onboarding")
        mark_step(a, "contract")
        flash("合同已生成，并连同 Start Letter 发给代理")
    else:
        flash("合同已生成，可在「代理合同」里下载检查")
    vault.put("agent", a)
    vault.audit("contract generated")
    return redirect(url_for("agent_detail", aid=aid))


@app.post("/agents/<aid>/contracts")
@login_required
def contract_upload(aid):
    """第 6 步：上传签回的合同。"""
    a = vault.get("agent", aid) or abort(404)
    f = save_upload("file") or abort(400)
    f.update(expiry=request.form.get("expiry") or None, kind="已签")
    a.setdefault("contracts", []).append(f)
    old_expiry = a.get("contract_expiry") or ""
    expiries = [c["expiry"] for c in a["contracts"] if c.get("expiry") and c.get("kind") in (None, "", "已签")]
    if expiries:
        a["contract_expiry"] = max(expiries)
    mark_step(a, "signed")
    if a.get("renewal") and (f.get("expiry") or "") > old_expiry:  # 签回的是续签合同
        a["renewal"]["signed"] = dt.date.today().isoformat()
        lv = a["renewal"].get("level")
        if lv in LEVELS:  # 签回续签合同后，按新等级计算佣金
            a["level"] = lv
            a["rates"] = [LEVELS[lv]["y1"], LEVELS[lv]["y2"], LEVELS[lv]["y3"]]
        a.setdefault("renewal_history", []).append(a.pop("renewal"))  # 下一次到期时重新开始
    if a.get("status") in ("潜在", "洽谈中"):
        a["status"] = "已签约"
    vault.put("agent", a)
    vault.audit("contract uploaded")
    flash("签回合同已加密保存")
    return redirect(url_for("agent_detail", aid=aid))


@app.post("/agents/<aid>/activate")
@login_required
def agent_activate(aid):
    """第 7 步：开通，发欢迎信，加入 News 名单。"""
    a = vault.get("agent", aid) or abort(404)
    a["code"] = a.get("code") or next_code()
    a["status"] = "活跃"
    if a.get("email") and request.form.get("send"):
        tpl = et.fill(et.WELCOME, contact=a.get("contact") or a["name"], code=a["code"])
        with mailer.Outbox(vault) as box:
            box.send(a["email"], tpl["subject"], tpl["body"], agent_id=aid, kind="onboarding")
    mark_step(a, "activate")
    vault.put("agent", a)
    flash(f"{a['name']} 已开通（{a['code']}），之后会自动收到每月 News")
    return redirect(url_for("agent_detail", aid=aid))


@app.route("/agents/<aid>/contracts/<int:i>")
@login_required
def contract_download(aid, i):
    a = vault.get("agent", aid) or abort(404)
    c = a["contracts"][i]
    if not c.get("fid"):
        abort(404)
    vault.audit("contract downloaded")
    return send_file(io.BytesIO(vault.get_file(c["fid"])), download_name=c["name"], as_attachment=True)


def _contract_data(aid, i):
    a = vault.get("agent", aid) or abort(404)
    if i >= len(a.get("contracts") or []):
        abort(404)
    c = a["contracts"][i]
    if not c.get("fid") or not c["name"].lower().endswith(".docx"):
        abort(404)
    return a, c, vault.get_file(c["fid"])


@app.route("/agents/<aid>/contracts/<int:i>/page/<int:n>.png")
@login_required
def contract_page(aid, i, n):
    """排版预览的第 n 页（图片在内存里生成，不落盘）。"""
    _, _, data = _contract_data(aid, i)
    pdf = contracts.to_pdf(data) or abort(404)
    if n >= contracts.page_count(pdf):
        abort(404)
    return send_file(io.BytesIO(contracts.page_png(pdf, n)), mimetype="image/png")


@app.route("/agents/<aid>/contracts/<int:i>/edit", methods=["GET", "POST"])
@login_required
def contract_edit(aid, i):
    """在线预览 Word 合同，直接改文字后保存（改动只影响改过的段落，其余格式不变）。"""
    a, c, data = _contract_data(aid, i)
    if request.method == "POST":
        try:
            edits = {int(k): str(v) for k, v in json.loads(request.form.get("edits") or "{}").items()}
        except (ValueError, AttributeError):
            abort(400)
        if edits:
            trash_contract(a, dict(c), "在线修改前的版本")
            c["fid"] = vault.put_file(contracts.apply_edits(data, edits))
            c["edited"] = dt.date.today().isoformat()
            vault.put("agent", a)
            vault.audit("contract edited online")
            flash(f"已保存 {len(edits)} 处修改")
        else:
            flash("没有改动")
        return redirect(url_for("contract_edit", aid=aid, i=i, back=safe_back(request.form.get("back")), mode="view"))
    pdf = contracts.to_pdf(data)
    return render_template("contract_edit.html", a=a, c=c, i=i, doc=contracts.outline(data),
                           pages=contracts.page_count(pdf) if pdf else 0, mode=request.args.get("mode") or "view",
                           back=safe_back(request.args.get("back")) or url_for("agent_detail", aid=aid),
                           is_renewal=(a.get("renewal") or {}).get("contract") == i)


# ---------- 佣金结算 ----------
@app.route("/commissions")
@login_required
def commissions():
    ts = terms()
    term = request.args.get("term") or (ts[-1] if ts else "")
    agents = agents_by_id()
    comms = vault.all("commission")
    rows = sorted([enrich(c, agents, comms) for c in comms if c["term"] == term],
                  key=lambda r: (STAGES.index(r["stage"]) if r["stage"] in STAGES else 9, r["agent"].get("code") or ""))
    return render_template("commissions.html", terms=ts, term=term, rows=rows, suggest=next_terms(ts),
                           total=sum(r["amount"] for r in rows), fees=sum(r["fees"] for r in rows))


@app.post("/commissions/terms")
@login_required
def term_create():
    name = re.sub(r"\s+", " ", request.form.get("name") or "").strip()[:60]
    if not name:
        flash("请填学期名称")
        return redirect(url_for("commissions"))
    ts = terms()
    if name not in ts:
        orders = [c.get("term_order", 0) for c in vault.all("commission")] + [t.get("order", 0) for t in vault.all("term")]
        vault.put("term", {"name": name, "order": max(orders, default=0) + 1, "created": dt.date.today().isoformat()})
        flash(f"已新建学期「{name}」，可以开始给代理建结算单")
    return redirect(url_for("commissions", term=name))


@app.post("/commissions/terms/delete")
@login_required
def term_delete():
    name = request.form.get("name") or ""
    if any(c["term"] == name for c in vault.all("commission")):
        flash("这个学期已经有结算单，不能删除")
    else:
        for t in vault.all("term"):
            if t["name"] == name:
                vault.delete("term", t["id"])
        flash(f"已删除空学期「{name}」")
    return redirect(url_for("commissions"))


def statement_form(c, f):
    c["agent_id"], c["term"] = f["agent_id"], f["term"].strip()
    orders = {t: i for i, t in enumerate(terms())}
    c["term_order"] = orders.get(c["term"], len(orders) + 100)
    students = []
    for i, name in enumerate(f.getlist("s_name")):
        tuition = fnum(f.getlist("s_tuition")[i])
        if not name.strip() or tuition is None:
            continue
        students.append({"name": name.strip(), "program": f.getlist("s_program")[i],
                         "courses": f.getlist("s_courses")[i].strip(), "origin": f.getlist("s_origin")[i],
                         "tuition": tuition, "note": f.getlist("s_note")[i].strip(),
                         "rate": fnum(f.getlist("s_rate")[i])})
    c["students"] = students
    for k in ("other_amount",):
        c[k] = fnum(f.get(k))
    for k in ("other_note", "notes"):
        c[k] = f.get(k, "").strip()
    if c.get("imported") and students:
        c["imported"] = False  # 补录了学生明细后，以系统计算为准
    return c


@app.route("/commissions/new", methods=["GET", "POST"])
@app.route("/commissions/<cid>", methods=["GET", "POST"])
@login_required
def commission_edit(cid=None):
    c = vault.get("commission", cid) if cid else {"term": request.args.get("term") or (terms() or [""])[-1],
                                                   "status": "草稿", "agent_id": request.args.get("agent"),
                                                   "students": []}
    if c is None:
        abort(404)
    agents = sorted(vault.all("agent"), key=lambda a: a.get("code") or "J9")
    if request.method == "POST":
        if request.form.get("delete"):
            vault.delete("commission", cid)
            return redirect(url_for("commissions", term=c["term"]))
        if c.get("locked"):
            flash("结算单已发出，金额已锁定。如需修改，请先「退回草稿」。")
            return redirect(url_for("commission_edit", cid=cid))
        c = vault.put("commission", statement_form(c, request.form))
        flash("已保存，佣金已自动计算")
        return redirect(url_for("commission_edit", cid=c["id"]))
    view = enrich(c, agents_by_id()) if c.get("agent_id") else dict(c, lines=[], calc=0, stage="草稿", scheme=None)
    inbound = []
    if c.get("agent_id"):
        inbound = [(m, f) for m in vault.all("mail") if m["direction"] == "in" and m.get("agent_id") == c["agent_id"]
                   for f in m.get("attachments", [])]
    return render_template("commission_form.html", c=view, agents=agents, terms=terms(), inbound=inbound,
                           cfg=settings())


@app.post("/commissions/<cid>/import-roster")
@login_required
def commission_import_roster(cid):
    c = vault.get("commission", cid) or abort(404)
    f = request.files.get("file") or abort(400)
    sheets = statements.parse_roster(f.read())
    sheet = request.form.get("sheet") or c["term"]
    pick = sheets.get(sheet) or (list(sheets.values())[-1] if sheets else [])
    if not pick:
        flash("没有在 Excel 里找到 Student Name 开头的学生名单")
    else:
        c["students"] = pick
        c["imported"] = False
        vault.put("commission", c)
        flash(f"已导入 {len(pick)} 名学生（工作表：{sheet if sheet in sheets else list(sheets)[-1]}）")
    return redirect(url_for("commission_edit", cid=cid))


def _statement_file(c, agents):
    e = enrich(c, agents)
    data = statements.statement_xlsx(e["agent"], c, e["lines"], e["calc"])
    name = f"JA Commission Statement {e['agent'].get('code', '')} {c['term']}.xlsx"
    return e, name, data


@app.route("/commissions/<cid>/statement.xlsx")
@login_required
def commission_statement(cid):
    c = vault.get("commission", cid) or abort(404)
    _, name, data = _statement_file(c, agents_by_id())
    return send_file(io.BytesIO(data), download_name=name, as_attachment=True)


def term_phrase(term):
    """'Sep-Dec Fall 26-27' -> '26-27 Sep-Dec Fall School Semester'（和以前发给代理的邮件标题一样）。"""
    m = re.match(r"^(.*?)\s*(\d{2}-\d{2})$", term or "")
    return f"{m.group(2)} {m.group(1)} School Semester" if m and m.group(1) else term


def roster_name(a, today=None):
    """附件名照以前的写法：<代理> Student Roster Update Nov-1-2026.xlsx"""
    today = today or dt.date.today()
    r = a.get("roster") or {}
    prefix = r.get("prefix") or re.sub(r"(\s+(last\s+)?update)?(\s+[A-Za-z]{3,9}-\d{1,2}-\d{2,4})?\.xlsx$", "",
                                        r.get("name") or "", flags=re.I).strip() or f"{a.get('name', '')} Student Roster"
    return f"{prefix} Update {today.strftime('%b')}-{today.day}-{today.year}.xlsx"


def roster_attachment(c, agents):
    """代理以前的 Student Roster + 本学期新工作表。代理还没有 Roster 底稿时，新建一个只有本学期的。"""
    e = enrich(c, agents)
    a = e["agent"]
    r = a.get("roster") or {}
    base = vault.get_file(r["fid"]) if r.get("fid") else None
    label = r.get("agent_label") or statements.last_agent_label(base) or a.get("contact") or a.get("name", "")
    data = statements.roster_workbook(base, c["term"], c, e["lines"], e["scheme"], label)
    return e, roster_name(a), data


def statement_mail(c, a, e):
    """结算邮件草稿：存过的（预览页里改过的）优先，否则按模板生成。"""
    d = c.get("draft_mail") or {}
    tpl = et.fill(et.COMMISSION_ROSTER if (a.get("roster") or c.get("draft_mail")) else et.COMMISSION_STATEMENT,
                  agent=d.get("greet") or a.get("greet") or a.get("name"), term=term_phrase(c["term"]),
                  amount=money(e["calc"]))
    return {"subject": d.get("subject") or tpl["subject"], "body": d.get("body") or tpl["body"],
            "to": d.get("to") or a.get("email") or ""}


def send_statement(c, agents, box, sender=None):
    """第 1 步：锁定金额，发结算邮件给代理。附件是代理的 Student Roster（以前所有学期 + 本学期新工作表）。"""
    e, name, data = roster_attachment(c, agents)
    a = e["agent"]
    m = statement_mail(c, a, e)
    box.send(m["to"], m["subject"], m["body"], [(name, data)], agent_id=a.get("id"), kind="statement",
             sender=sender, commission_id=c["id"])
    c["locked"] = {"lines": e["lines"], "total": e["calc"], "y1_rate": y1_rate_used(e["scheme"], ay_group(c), c)}
    for prior in ay_group(c):  # 本单已包含补差，之前的单子不再重复补
        if prior["id"] != c["id"] and prior.get("locked") and any(l.get("kind") == "TOPUP" for l in e["lines"]):
            prior["topped_up_by"] = c["id"]
            vault.put("commission", prior)
    c["status"] = "已发结算单"
    c["sent_on"] = dt.date.today().isoformat()
    c["sent_file"] = {"name": name, "fid": vault.put_file(data)}
    vault.put("commission", c)
    # 发出去的这份就是代理最新的 Roster，下个学期在它后面再加工作表
    a = vault.get("agent", a["id"])
    old = a.get("roster") or {}
    a["roster"] = dict(old, name=name, fid=c["sent_file"]["fid"], updated=c["sent_on"])
    vault.put("agent", a)


@app.post("/commissions/<cid>/send")
@login_required
def commission_send(cid):
    c = vault.get("commission", cid) or abort(404)
    agents = agents_by_id()
    if c.get("locked"):
        flash("这张结算单已经发出过了")
        return redirect(url_for("commission_edit", cid=cid))
    to = (c.get("draft_mail") or {}).get("to") or agents.get(c["agent_id"], {}).get("email")
    if not to:
        flash("这个代理还没有邮箱，先在预览页填收件人，或到代理资料里补上")
        return redirect(url_for("commission_mail", cid=cid))
    with mailer.Outbox(vault) as box:
        send_statement(c, agents, box, sender=request.form.get("sender") or None)
    flash("结算邮件和 Student Roster 已发给代理，等代理确认并回发票")
    return redirect(url_for("commission_edit", cid=cid))


@app.route("/commissions/<cid>/mail", methods=["GET", "POST"])
@login_required
def commission_mail(cid):
    """发出前预览：邮件正文可以改，附件（Student Roster）每个工作表都能直接看。"""
    c = vault.get("commission", cid) or abort(404)
    agents = agents_by_id()
    a = agents.get(c["agent_id"]) or abort(404)
    if request.method == "POST" and not c.get("locked"):
        if request.form.get("roster_reset"):
            a.pop("roster", None)
            vault.put("agent", a)
            flash("已去掉 Roster 底稿，附件改为只有本学期的新表")
        elif request.files.get("roster") and request.files["roster"].filename:
            f = request.files["roster"]
            if not f.filename.lower().endswith(".xlsx"):
                flash("请上传 .xlsx")
            else:
                data = f.read()
                old = a.get("roster") or {}
                a["roster"] = {"name": os.path.basename(f.filename), "fid": vault.put_file(data),
                               "updated": dt.date.today().isoformat(), "agent_label": old.get("agent_label", "")}
                vault.put("agent", a)
                flash("已换成新的 Roster 底稿，附件已按它重新生成")
        else:
            d = c.get("draft_mail") or {}
            for k in ("to", "subject", "body"):
                d[k] = (request.form.get(k) or "").strip()
            c["draft_mail"] = d
            if request.form.get("agent_label") is not None:
                r = a.get("roster") or {}
                r["agent_label"] = request.form.get("agent_label", "").strip()
                if a.get("roster"):
                    a["roster"] = r
                    vault.put("agent", a)
            vault.put("commission", c)
            flash("邮件草稿已保存")
        return redirect(url_for("commission_mail", cid=cid))
    if c.get("locked") and c.get("sent_file"):
        e = enrich(c, agents)
        name, data = c["sent_file"]["name"], vault.get_file(c["sent_file"]["fid"])
    else:
        e, name, data = roster_attachment(c, agents)
    m = statement_mail(c, a, e)
    sheets = statements.preview(data)
    r = a.get("roster") or {}
    return render_template("commission_mail.html", c=e, a=a, m=m, att_name=name, sheets=sheets,
                           new_sheet=statements.sheet_title(c["term"]), roster=r,
                           agent_label=r.get("agent_label") or statements.last_agent_label(
                               vault.get_file(r["fid"]) if r.get("fid") else None) or a.get("name", ""))


@app.route("/commissions/<cid>/roster.xlsx")
@login_required
def commission_roster(cid):
    c = vault.get("commission", cid) or abort(404)
    if c.get("locked") and c.get("sent_file"):
        name, data = c["sent_file"]["name"], vault.get_file(c["sent_file"]["fid"])
    else:
        _, name, data = roster_attachment(c, agents_by_id())
    return send_file(io.BytesIO(data), download_name=name, as_attachment=True)


@app.post("/commissions/send-all")
@login_required
def commission_send_all():
    term = request.form["term"]
    agents = agents_by_id()
    drafts = [c for c in vault.all("commission") if c["term"] == term and stage_of(c) == "草稿"
              and not c.get("imported") and agents.get(c["agent_id"], {}).get("email")]
    with mailer.Outbox(vault) as box:
        for c in drafts:
            send_statement(c, agents, box)
    flash(f"已给 {len(drafts)} 个代理发出结算单")
    return redirect(url_for("commissions", term=term))


@app.post("/commissions/<cid>/confirm")
@login_required
def commission_confirm(cid):
    """第 2 步：代理确认，登记发票。"""
    c = vault.get("commission", cid) or abort(404)
    inv = save_upload("invoice")
    if not inv and request.form.get("from_mail"):
        name, fid = request.form["from_mail"].split("|", 1)
        inv = {"name": name, "fid": fid, "uploaded": dt.date.today().isoformat()}
    c["invoice"] = inv
    c["confirmed_on"] = dt.date.today().isoformat()
    c["status"] = "代理已确认"
    vault.put("commission", c)
    flash("已登记代理确认" + ("和发票" if inv else "（未附发票）"))
    return redirect(url_for("commission_edit", cid=cid))


@app.post("/commissions/<cid>/accountant")
@login_required
def commission_accountant(cid):
    """第 3 步：把付款需要的信息（只有这些）发给会计。"""
    c = vault.get("commission", cid) or abort(404)
    cfg = settings()
    if not cfg.get("accountant_email"):
        flash("请先在「设置与安全」里填会计邮箱")
        return redirect(url_for("commission_edit", cid=cid))
    agents = agents_by_id()
    e, name, data = _statement_file(c, agents)
    a = e["agent"]
    method = request.form.get("method") or ("E-Transfer" if a.get("etransfer") else "Cheque")
    tpl = et.fill(et.ACCOUNTANT, agent=a.get("payee") or a.get("name"), code=a.get("code"), term=c["term"],
                  amount=money(e["amount"]), method=method, payee=a.get("payee") or a.get("name"),
                  etransfer=a.get("etransfer") or "-", bank=a.get("detail") or "-")
    files = [(name, data)]
    if c.get("invoice"):
        files.append((c["invoice"]["name"], vault.get_file(c["invoice"]["fid"])))
    with mailer.Outbox(vault) as box:
        box.send(cfg["accountant_email"], tpl["subject"], tpl["body"], files, commission_id=cid, kind="accountant")
    c.update(status="已交会计", to_accountant_on=dt.date.today().isoformat(), pay_method=method)
    vault.put("commission", c)
    flash("已发给会计付款")
    return redirect(url_for("commission_edit", cid=cid))


@app.post("/commissions/<cid>/paid")
@login_required
def commission_paid(cid):
    """第 4 步：登记付款凭证（支票 / E-Transfer 截图），可通知代理。"""
    c = vault.get("commission", cid) or abort(404)
    c["proof"] = save_upload("proof")
    c.update(status="已付", paid_on=request.form.get("paid_on") or dt.date.today().isoformat(),
             pay_method=request.form.get("method") or c.get("pay_method"), pay_ref=request.form.get("ref", ""))
    vault.put("commission", c)
    if request.form.get("notify"):
        e = enrich(c, agents_by_id())
        a = e["agent"]
        if a.get("email"):
            tpl = et.fill(et.PAYMENT_NOTICE, agent=a.get("name"), term=c["term"], amount=money(e["amount"]),
                          method=c["pay_method"], date=c["paid_on"])
            with mailer.Outbox(vault) as box:
                box.send(a["email"], tpl["subject"], tpl["body"], agent_id=a["id"], commission_id=cid, kind="statement")
    flash("已登记付款，本单结算完成")
    return redirect(url_for("commission_edit", cid=cid))


@app.post("/commissions/<cid>/reopen")
@login_required
def commission_reopen(cid):
    c = vault.get("commission", cid) or abort(404)
    c.pop("locked", None)
    c["status"] = "草稿"
    vault.put("commission", c)
    flash("已退回草稿，可以修改后重新发送")
    return redirect(url_for("commission_edit", cid=cid))


@app.route("/commissions/export")
@login_required
def commissions_export():
    term = request.args["term"]
    agents = agents_by_id()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\\/*?:\[\]]", "-", term)[:31]
    ws.append(["代理名称", "代理编号", "佣金方案", "学生数", "已缴学费总额", "Commission 金额", "状态", "发结算单", "代理确认",
               "交会计", "付款日期", "付款方式", "E-Transfer", "支票抬头"])
    comms = vault.all("commission")
    for r in sorted([enrich(c, agents, comms) for c in comms if c["term"] == term],
                    key=lambda r: r["agent"].get("code") or ""):
        a = r["agent"]
        ws.append([a.get("name"), a.get("code"), scheme_label(a), len(r.get("students", [])), r["fees"], r["amount"],
                   r["stage"], r.get("sent_on"), r.get("confirmed_on"), r.get("to_accountant_on"), r.get("paid_on"),
                   r.get("pay_method"), a.get("etransfer"), a.get("payee")])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    vault.audit("commission export")
    return send_file(buf, download_name=f"Commission {ws.title}.xlsx", as_attachment=True)


# ---------- 开发新代理 ----------
def needs_followup(p):
    if p.get("stage") not in ("已发邮件", "已跟进") or p.get("unsubscribed"):
        return False
    last = p.get("last_email_on")
    return bool(last) and (dt.date.today() - dt.date.fromisoformat(last)).days >= 7


@app.route("/prospects")
@login_required
def prospects():
    f = {k: request.args.get(k, "") for k in ("country", "tier", "stage", "q")}
    rows = vault.all("prospect")
    countries = sorted({p.get("country", "") for p in rows if p.get("country")})
    if f["country"]:
        rows = [p for p in rows if p.get("country") == f["country"]]
    if f["tier"]:
        rows = [p for p in rows if p.get("tier") == f["tier"]]
    if f["stage"]:
        rows = [p for p in rows if (p.get("stage") or "未联系") == f["stage"]]
    if f["q"]:
        q = f["q"].lower()
        rows = [p for p in rows if q in " ".join(str(p.get(k, "")) for k in ("agency", "contact", "city", "email")).lower()]
    rows.sort(key=lambda p: ({"A": 0, "B": 1, "C": 2}.get(p.get("tier"), 3), p.get("country", ""), p.get("agency", "")))
    allp = vault.all("prospect")
    jobs = sorted(vault.all("ai_job"), key=lambda j: j["started"], reverse=True)[:5]
    return render_template("prospects.html", rows=rows, f=f, countries=countries, jobs=jobs,
                           review=[p for p in allp if p.get("stage") == "待审核"],
                           counts={s: sum(1 for p in allp if (p.get("stage") or "未联系") == s) for s in PROSPECT_STAGES},
                           followups=[p for p in allp if needs_followup(p)], regions=i18n.TARGET_REGIONS,
                           has_ai=ai.available(settings()), ai_mode=ai.mode(settings()))


def import_tracker(data: bytes) -> int:
    """Agent Tracker Excel → 潜在代理（已存在的机构跳过）。网页导入和加密收件箱共用。"""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    existing = {(p.get("agency") or "").lower() for p in vault.all("prospect")}
    n = 0

    def val(r, idx, *names):
        for nm in names:
            if nm in idx and r[idx[nm]] not in (None, ""):
                v = str(r[idx[nm]]).strip()
                if not v.lower().startswith(("not publicly", "not separately", "to confirm")):
                    return v
        return ""

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows or "Agency" not in [str(h).strip() for h in rows[0] if h]:
            continue
        idx = {str(h).strip(): i for i, h in enumerate(rows[0]) if h}
        for r in rows[1:]:
            agency = val(r, idx, "Agency")
            if not agency or agency.lower() in existing:
                continue
            existing.add(agency.lower())
            grade = val(r, idx, "Lead Grade", "Research Completeness")[:1]
            vault.put("prospect", {
                "agency": agency, "country": val(r, idx, "Country"),
                "city": val(r, idx, "City / Office", "Office / City", "Office / Address"),
                "contact": val(r, idx, "Contact Person", "Contact Name"), "role": val(r, idx, "Role", "Title / Role"),
                "email": val(r, idx, "Partnership Email", "Direct Email", "Email", "General Email"),
                "phone": val(r, idx, "Phone", "Phone / WhatsApp"), "whatsapp": val(r, idx, "WhatsApp"),
                "website": val(r, idx, "Website"), "tier": grade if grade in "ABC" else "",
                "priority": val(r, idx, "Priority"), "k12_fit": val(r, idx, "G8-12 Fit", "High School / G8-12 Fit",
                                                                      "G8-12 / High School Fit"),
                "notes": val(r, idx, "Notes", "Verification / Notes", "Verification Notes", "Recommended First Contact"),
                "source_url": val(r, idx, "Source URL"), "stage": "未联系", "source": "Agent Tracker 导入"})
            n += 1
    return n


@app.post("/prospects/import")
@login_required
def prospects_import():
    f = request.files.get("file") or abort(400)
    n = import_tracker(f.read())
    flash(f"从 Tracker 导入 {n} 个潜在代理（已存在的自动跳过）")
    return redirect(url_for("prospects"))


def _run_search(job_id, cfg, countries, n, extra):
    job = vault.get("ai_job", job_id)
    exclude = {p.get("agency", "") for p in vault.all("prospect")} | {a.get("name", "") for a in vault.all("agent")}
    found = 0

    def progress(msg):
        job["progress"] = msg
        vault.put("ai_job", job)

    try:
        for country in countries:
            progress(f"正在搜索 {i18n.COUNTRY_ZH.get(country, country)}…")
            for lead in ai.search_agencies(cfg, country, n, extra, exclude, i18n.lang_for(country), progress):
                vault.put("prospect", {**lead, "tier": lead.pop("size_tier", ""), "stage": "待审核",
                                       "source": f"AI 搜索（{ {'local': '本地', 'hybrid': '本地 + GPT'}.get(ai.mode(cfg), 'Claude')}）"
                                                 f" {dt.date.today().isoformat()}"})
                exclude.add(lead["agency"])
                found += 1
        job.update(status="完成", progress=f"找到 {found} 个，等待你审核")
    except Exception as e:  # 失败原因显示在页面上
        job.update(status="失败", progress=str(e)[:300])
    vault.put("ai_job", job)


@app.post("/prospects/search")
@login_required
def prospects_search():
    cfg = settings()
    if not ai.available(cfg):
        flash("请先在「设置与安全」里设置 AI：本地模型地址 + 搜索地址，或 Claude API Key")
        return redirect(url_for("prospects"))
    countries = request.form.getlist("countries")
    if not countries:
        flash("请至少选一个国家")
        return redirect(url_for("prospects"))
    n = min(int(request.form.get("n") or 10), 25)
    job = vault.put("ai_job", {"countries": countries, "n": n, "status": "进行中", "progress": "开始搜索…",
                               "started": dt.datetime.now().isoformat(timespec="minutes")})
    threading.Thread(target=_run_search, args=(job["id"], cfg, countries, n, request.form.get("extra", "")),
                     daemon=True).start()
    flash("AI 正在联网搜索，每个国家约需 2–10 分钟，完成后结果会出现在「待审核」里")
    return redirect(url_for("prospects"))


@app.post("/prospects/review")
@login_required
def prospects_review():
    keep = set(request.form.getlist("keep"))
    for p in vault.all("prospect"):
        if p.get("stage") == "待审核" and p["id"] in set(request.form.getlist("ids")):
            if p["id"] in keep:
                p["stage"] = "未联系"
                vault.put("prospect", p)
            else:
                vault.delete("prospect", p["id"])
    flash("审核完成")
    return redirect(url_for("prospects"))


@app.route("/prospects/<pid>", methods=["GET", "POST"])
@login_required
def prospect_detail(pid):
    p = vault.get("prospect", pid) or abort(404)
    if request.method == "POST":
        for k in ("agency", "country", "city", "contact", "role", "email", "phone", "whatsapp", "website", "tier",
                  "stage", "notes"):
            if k in request.form:
                p[k] = request.form[k].strip()
        vault.put("prospect", p)
        flash("已保存")
        return redirect(url_for("prospect_detail", pid=pid))
    mails = sorted([m for m in vault.all("mail") if m.get("prospect_id") == pid], key=lambda m: m["at"], reverse=True)
    return render_template("prospect_detail.html", p=p, mails=mails, lang=i18n.lang_for(p.get("country")))


@app.post("/prospects/<pid>/convert")
@login_required
def prospect_convert(pid):
    """潜在代理谈成了：转成代理，入驻流程从第 3 步继续。"""
    p = vault.get("prospect", pid) or abort(404)
    a = vault.put("agent", {"name": p["agency"], "contact": p.get("contact", ""), "email": p.get("email", ""),
                            "phone": p.get("phone", ""), "country": p.get("country", ""), "city": p.get("city", ""),
                            "website": p.get("website", ""), "status": "洽谈中", "level": "standard",
                            "rates": [20, 10, 5], "source": "开发新代理", "contracts": [],
                            "onboarding": {"contact": p.get("first_email_on") or dt.date.today().isoformat(),
                                           "meeting": dt.date.today().isoformat()}})
    p["stage"] = "已转为代理"
    p["agent_id"] = a["id"]
    vault.put("prospect", p)
    for m in vault.all("mail"):
        if m.get("prospect_id") == pid:
            m["agent_id"] = a["id"]
            vault.put("mail", m)
    flash("已转为代理，接下来：申请表与背景核实")
    return redirect(url_for("agent_detail", aid=a["id"]))


@app.route("/prospects/outreach", methods=["GET", "POST"])
@login_required
def outreach():
    ids = dict.fromkeys(request.form.getlist("ids") if request.method == "POST" else request.args.getlist("ids"))
    picked = [p for p in (vault.get("prospect", i) for i in ids) if p]
    if request.method == "POST" and request.form.get("go"):
        english = {"subject": request.form["subject"], "body": request.form["body"]}
        custom = {k[3:]: {"label": i18n.TRANSLATIONS.get(k[3:], {}).get("label", k[3:]), "subject": "",
                          "body": v} for k, v in request.form.items() if k.startswith("tr_") and v.strip()}
        sent = skipped = 0
        follow = request.form.get("kind") == "followup"
        t = vault.get("template", request.form.get("tpl", ""))
        if t and request.form.get("save_tpl"):
            t.update(subject=english["subject"], body=english["body"])
            t.setdefault("translations", {}).update({k: v["body"] for k, v in custom.items()})
            vault.put("template", t)
        with mailer.Outbox(vault) as box:
            for p in picked:
                if not p.get("email") or p.get("unsubscribed"):
                    skipped += 1
                    continue
                lang = i18n.lang_for(p.get("country"))
                subject, body, _ = i18n.compose(p, english, custom.get(lang, False))
                box.send(p["email"], subject, body, prospect_id=p["id"], kind="outreach")
                today = dt.date.today().isoformat()
                p["first_email_on"] = p.get("first_email_on") or today
                p["last_email_on"] = today
                p["emails_sent"] = p.get("emails_sent", 0) + 1
                p["stage"] = "已跟进" if follow or p.get("stage") in ("已发邮件", "已跟进") else "已发邮件"
                vault.put("prospect", p)
                sent += 1
        flash(f"已发出 {sent} 封{'跟进' if follow else '开发'}邮件" + (f"，{skipped} 个没有邮箱或已退订，已跳过" if skipped else ""))
        return redirect(url_for("prospects"))
    langs = sorted({i18n.lang_for(p.get("country")) for p in picked} - {None})
    follow = request.values.get("kind") == "followup"
    tpls = all_templates()
    mine = [t for t in tpls if t["kind"] == ("followup" if follow else "outreach")] or tpls
    tpl = next((t for t in tpls if t["id"] == request.args.get("tpl")), None) or \
        next((t for t in mine if "推荐" in t["name"]), mine[0])
    labels = {k: v["label"] for k, v in i18n.TRANSLATIONS.items()}
    return render_template("outreach.html", picked=picked, english=tpl, langs=langs, kind="followup" if follow else "",
                           tpl=tpl, tpls=tpls,
                           translations={l: {"label": labels.get(l, l), "body": tpl.get("translations", {}).get(l, "")}
                                         for l in langs},
                           has_ai=ai.available(settings()), ai_mode=ai.mode(settings()))


def all_templates():
    """邮件模板库；第一次打开时放入默认模板（新版开发邮件、原短版、跟进邮件）。"""
    ts = vault.all("template")
    if not ts:
        for t in i18n.DEFAULT_TEMPLATES:
            vault.put("template", json.loads(json.dumps(t)))
        ts = vault.all("template")
    return sorted(ts, key=lambda t: (t.get("kind") != "outreach", t.get("kind", ""), t["name"]))


@app.route("/mail/templates", methods=["GET", "POST"])
@app.route("/mail/templates/<tid>", methods=["GET", "POST"])
@login_required
def mail_templates(tid=None):
    tpls = all_templates()
    t = next((x for x in tpls if x["id"] == tid), None) if tid else None
    if request.method == "POST":
        if request.form.get("delete") and t:
            vault.delete("template", t["id"])
            flash("模板已删除")
            return redirect(url_for("mail_templates"))
        t = t or {"translations": {}}
        for k in ("name", "kind", "subject", "body"):
            t[k] = request.form.get(k, "").strip()
        t = vault.put("template", t)
        flash("模板已保存")
        return redirect(url_for("mail_templates", tid=t["id"]))
    return render_template("templates.html", tpls=tpls, t=t,
                           labels={k: v["label"] for k, v in i18n.TRANSLATIONS.items()})


@app.post("/prospects/translate")
@login_required
def prospects_translate():
    cfg = settings()
    if not ai.available(cfg):
        return {"error": "请先在设置里设置 AI（本地模型或 Claude API Key）"}, 400
    lang = request.form["lang"]
    try:
        out = ai.translate(cfg, request.form["subject"], request.form["body"],
                           i18n.TRANSLATIONS.get(lang, {}).get("label", lang))
    except Exception as e:
        return {"error": str(e)[:300]}, 500
    return out


# ---------- 邮件 ----------
@app.route("/mail")
@login_required
def mail():
    box = request.args.get("box", "in")
    agents = agents_by_id()
    msgs = sorted([m for m in vault.all("mail") if m["direction"] == box], key=lambda m: m["at"], reverse=True)
    return render_template("mail.html", box=box, msgs=msgs, agents=agents,
                           campaigns=sorted(vault.all("campaign"), key=lambda c: c.get("next_at") or "9"),
                           live=mailer.mail_on(settings()))


@app.route("/mail/compose", methods=["GET", "POST"])
@login_required
def compose():
    agents = sorted([a for a in vault.all("agent") if a.get("email")], key=lambda a: a.get("code") or "J9")
    if request.method == "POST":
        f = request.form
        c = {"subject": f["subject"], "body": f["body"], "audience": f["audience"],
             "agent_ids": f.getlist("agent_ids"), "cadence": f["cadence"], "sender": f.get("sender") or None}
        when = f.get("send_at")
        if f["cadence"] == "立即发送" and not when:
            sent = mailer.send_to_agents(vault, mailer.recipients(vault, c), c["subject"], c["body"],
                                         sender=c["sender"])
            flash(f"已处理 {len(sent)} 封邮件")
            return redirect(url_for("mail", box="out"))
        c.update(active=True, next_at=when or dt.datetime.now().isoformat(timespec="minutes"),
                 created=dt.date.today().isoformat())
        vault.put("campaign", c)
        flash("定时推送已创建")
        return redirect(url_for("mail"))
    return render_template("compose.html", agents=agents, pre=request.args.get("agent"),
                           kit=[x for x in all_assets() if x.get("url") and x.get("share")])


@app.post("/mail/campaign/<cid>")
@login_required
def campaign_toggle(cid):
    c = vault.get("campaign", cid) or abort(404)
    if request.form.get("delete"):
        vault.delete("campaign", cid)
    else:
        c["active"] = not c.get("active")
        vault.put("campaign", c)
    return redirect(url_for("mail"))


@app.post("/mail/fetch")
@login_required
def mail_fetch():
    try:
        flash(f"收到 {mailer.fetch_inbox(vault)} 封新邮件")
    except Exception as e:
        flash(f"收信失败：{e}")
    return redirect(url_for("mail"))


@app.route("/mail/msg/<mid>")
@login_required
def mail_msg(mid):
    m = vault.get("mail", mid) or abort(404)
    if m.get("status") == "未读":
        m["status"] = "已读"
        vault.put("mail", m)
    return render_template("mail_msg.html", m=m, a=vault.get("agent", m.get("agent_id")) if m.get("agent_id") else None,
                           p=vault.get("prospect", m.get("prospect_id")) if m.get("prospect_id") else None)


# ---------- 学生申请进度 ----------


def student_from_form(st, f):
    for k in ("name", "agent_id", "program", "grade", "intake", "country", "courses", "notes", "origin"):
        if k in f:
            st[k] = f.get(k, "").strip()
    st["tuition"] = fnum(f.get("tuition"))
    new = f.get("stage")
    if new in STUDENT_STAGES and new != st.get("stage"):
        st["stage"] = new
        st.setdefault("history", {})[new] = f.get("stage_date") or dt.date.today().isoformat()
    return st


_adm_last = {"at": 0, "result": None, "error": None}


def sync_admissions(force=False):
    """同步官网入学申请（最多每 10 分钟一次，除非手动点「立即同步」）。"""
    if not (admissions.enabled() or ap_courses.enabled()) or (not force and time.time() - _adm_last["at"] < 600):
        return _adm_last
    _adm_last["at"] = time.time()
    if admissions.enabled():
        try:
            _adm_last.update(result=admissions.sync(vault, STUDENT_STAGES), error=None)
            if _adm_last["result"]["new"]:
                vault.audit(f"admissions sync: {_adm_last['result']['new']} new")
        except Exception as e:  # noqa: BLE001  官网数据库连不上不影响代理系统
            _adm_last["error"] = f"{type(e).__name__}: {e}"[:200]
    if ap_courses.enabled():
        try:
            r = ap_courses.sync(vault)
            if r["new"]:
                vault.audit(f"AP courses sync: {r['new']} new")
        except Exception as e:  # noqa: BLE001  AP 系统连不上不影响代理系统
            app.logger.warning("AP courses sync failed: %s", e)
    return _adm_last


@app.post("/students/sync")
@login_required
def students_sync():
    r = sync_admissions(force=True)
    if r["error"]:
        flash(f"同步官网申请失败：{r['error']}")
    elif r["result"]:
        x = r["result"]
        flash(f"已同步官网申请：新增 {x['new']} 份，更新 {x['updated']} 份；待分配代理 {x['unassigned']} 份")
    return redirect(url_for("students", **({"stage": "待分配"} if r.get("result") and r["result"]["unassigned"] else {})))


@app.post("/students/<sid>/assign")
@login_required
def student_assign(sid):
    st = vault.get("student", sid) or abort(404)
    aid = request.form.get("agent_id", "")
    if aid == "self":
        aid = ""
    if aid and not vault.get("agent", aid):
        abort(400)
    st.update(agent_id=aid, agent_manual=True, agent_auto="")
    if not aid:
        st["uses_agent"] = False  # 选「学生自己申请」= 确认没有代理，不再出现在待分配里
    vault.put("student", st)
    flash(f"{st.get('name', '')}：已分配给 {vault.get('agent', aid)['name']}" if aid else f"{st.get('name', '')}：标为学生自己申请")
    back = request.form.get("back") or ""
    return redirect(back if back.startswith("/") and not back.startswith("//") else url_for("students"))


@app.route("/students")
@login_required
def students():
    sync_admissions()
    f = {k: request.args.get(k, "") for k in ("stage", "agent", "intake", "q")}
    agents = agents_by_id()
    rows = vault.all("student")
    unassigned = [s for s in rows if (s.get("web_no") or s.get("agent_hint")) and s.get("uses_agent") and not s.get("agent_id")]
    intakes = sorted({s.get("intake") for s in rows if s.get("intake")})
    counts = {g: sum(1 for s in rows if s.get("stage") == g) for g in STUDENT_STAGES}
    if f["stage"] == "待分配":
        rows = unassigned
    elif f["stage"]:
        rows = [s for s in rows if s.get("stage") == f["stage"]]
    if f["agent"]:
        rows = [s for s in rows if s.get("agent_id") == f["agent"]]
    if f["intake"]:
        rows = [s for s in rows if s.get("intake") == f["intake"]]
    if f["q"]:
        q = f["q"].lower()
        rows = [s for s in rows if q in (s.get("name", "") + " " + agents.get(s.get("agent_id"), {}).get("name", "")).lower()]
    rows.sort(key=lambda s: (STUDENT_STAGES.index(s["stage"]) if s.get("stage") in STUDENT_STAGES else 9,
                             max((s.get("history") or {"": ""}).values())), reverse=False)
    return render_template("students.html", rows=rows, agents=agents, counts=counts, f=f, intakes=intakes,
                           unassigned=len(unassigned), adm_on=admissions.enabled(), adm=_adm_last,
                           agent_list=sorted(agents.values(), key=lambda a: a.get("name", "")))


@app.route("/students/new", methods=["GET", "POST"])
@app.route("/students/<sid>", methods=["GET", "POST"])
@login_required
def student_edit(sid=None):
    st = vault.get("student", sid) if sid else {"stage": "", "agent_id": request.args.get("agent", ""),
                                                 "program": "FT1", "intake": (terms() or [""])[-1], "history": {}}
    if st is None:
        abort(404)
    if request.method == "POST":
        if request.form.get("delete"):
            vault.delete("student", sid)
            flash("已删除")
            return redirect(url_for("students"))
        if not st.get("stage") and not request.form.get("stage"):
            st["stage"] = "收到申请"
            st.setdefault("history", {})["收到申请"] = dt.date.today().isoformat()
        before = st.get("agent_id")
        st = student_from_form(st, request.form)
        if st.get("web_no") and st.get("agent_id") != before:
            st.update(agent_manual=True, agent_auto="")
        doc = save_upload("doc")
        if doc:
            st.setdefault("docs", []).append(doc)
        st = vault.put("student", st)
        flash("已保存")
        return redirect(url_for("student_edit", sid=st["id"]))
    agents = sorted(vault.all("agent"), key=lambda a: a.get("name", ""))
    return render_template("student_form.html", st=st, agents=agents, terms=terms(), agent=agents_by_id().get(st.get("agent_id")))


@app.post("/students/<sid>/stage")
@login_required
def student_stage(sid):
    st = vault.get("student", sid) or abort(404)
    new = request.form.get("stage")
    if new in STUDENT_STAGES:
        st["stage"] = new
        st.setdefault("history", {})[new] = dt.date.today().isoformat()
        vault.put("student", st)
        flash(f"{st['name']}：{new}")
    back = request.form.get("back") or ""
    return redirect(back if back.startswith("/") and not back.startswith("//") else url_for("student_edit", sid=sid))


@app.post("/students/<sid>/to-statement")
@login_required
def student_to_statement(sid):
    """已入学的学生一键加进该代理这个学期的结算单（没有就新建草稿）。"""
    st = vault.get("student", sid) or abort(404)
    if not st.get("agent_id") or not st.get("intake"):
        flash("先填写代理和入学学期")
        return redirect(url_for("student_edit", sid=sid))
    c = next((c for c in vault.all("commission") if c["agent_id"] == st["agent_id"] and c["term"] == st["intake"]), None)
    if c and c.get("locked"):
        flash("这个学期的结算单已经发出，金额已锁定。请先在结算单里「退回草稿」再加。")
        return redirect(url_for("commission_edit", cid=c["id"]))
    if not c:
        orders = {t: i for i, t in enumerate(terms())}
        c = {"agent_id": st["agent_id"], "term": st["intake"], "term_order": orders.get(st["intake"], len(orders) + 100),
             "status": "草稿", "students": []}
    row = {"name": st["name"], "program": st.get("program") or "FT1", "courses": st.get("courses") or "",
           "origin": st.get("origin") or "International", "tuition": st.get("tuition") or 0, "note": "", "rate": None,
           "student_id": st["id"]}
    c["students"] = [r for r in c.get("students", []) if r.get("student_id") != st["id"]] + [row]
    c = vault.put("commission", c)
    st["commission_id"] = c["id"]
    vault.put("student", st)
    flash(f"已加入 {st['intake']} 结算单，佣金已自动计算")
    return redirect(url_for("commission_edit", cid=c["id"]))


# ---------- 合同续签 ----------
RENEW_DAYS = 120   # 到期前多少天开始出现在续签列表
AUTO_DRAFT_DAYS = 90  # 到期前多少天系统自动生成续签合同（只生成，不发送）


def current_ay(today=None):
    """现在所在的学年：9 月起算新学年，例如 2026-10 → 26-27。"""
    today = today or dt.date.today()
    y = today.year % 100 if today.month >= 9 else today.year % 100 - 1
    return f"{y:02d}-{y + 1:02d}"


def last_ay(today=None):
    today = today or dt.date.today()
    y = today.year % 100 if today.month >= 9 else today.year % 100 - 1
    return f"{y - 1:02d}-{y:02d}"


def agent_ay_stats(aid, ay, comms=None, agents=None):
    comms = comms if comms is not None else vault.all("commission")
    agents = agents or agents_by_id()
    mine = [c for c in comms if c["agent_id"] == aid and academic_year(c["term"]) == ay]
    names = {(s.get("name") or "").lower() for c in mine for s in c.get("students", [])}
    ft1 = {(s.get("name") or "").lower() for c in mine for s in c.get("students", []) if s.get("program") == "FT1"}
    return {"ft1": len(ft1), "students": len(names), "commission": sum(enrich(c, agents, comms)["amount"] for c in mine),
            "fees": sum(total_fees(c) for c in mine), "terms": len(mine)}


def suggest_level(a, stats):
    """上一学年第一年学生满 6 人建议高级；旧合同第一年 25% 以上的，建议高级，避免续签后比例变低。"""
    if stats["ft1"] >= 6:
        return "premier"
    if a.get("level") in LEVELS:
        return a["level"]
    return "premier" if (a.get("rates") or [0])[0] >= LEVELS["premier"]["y1"] else "standard"


def renewal_period(a):
    old = a.get("contract_expiry")
    start = dt.date.fromisoformat(old) + dt.timedelta(days=1) if old and days_left(old) >= 0 else dt.date.today()
    return contracts.default_period(start)


def draft_renewal(a, level, start, end):
    fname, data = contracts.generate(a, level, start, end, renew=True)
    a.setdefault("contracts", []).append({"name": fname, "fid": vault.put_file(data), "expiry": end.isoformat(),
                                          "uploaded": dt.date.today().isoformat(), "kind": "续签待签",
                                          "tpl": TEMPLATE_REV})
    a["renewal"] = {"drafted": dt.date.today().isoformat(), "tpl": TEMPLATE_REV, "level": level, "start": start.isoformat(),
                    "end": end.isoformat(), "contract": len(a["contracts"]) - 1}
    return fname, data


def renewal_candidates():
    return sorted([a for a in vault.all("agent") if a.get("contract_expiry") and a.get("status") in ("活跃", "已签约")
                   and days_left(a["contract_expiry"]) <= RENEW_DAYS and not (a.get("renewal") or {}).get("signed")],
                  key=lambda a: a["contract_expiry"])


TEMPLATE_REV = 3  # 2：乙方信息按甲方的字体字号填写；3：签名行每页底部 + 页码，乙方地址电话等留空给代理填


def refresh_old_drafts():
    """清理旧模板生成的合同（POLO 2026-10-27：不要留有错误的旧版）：
    - 续签草稿：只留当前这一份，按新模板重新生成；被替换掉的、续签已签回后剩下的草稿都删掉。
    - 新签合同草稿（待签（系统生成））：每个代理只留最新一份，按新模板重新生成。
    代理签回上传的合同（已签）和导入的旧合同一律不动。邮件还没真正发出过（演示模式），所以重做不影响代理。"""
    for a in vault.all("agent"):
        cs = a.get("contracts") or []
        r = a.get("renewal") or {}
        cur = r.get("contract") if r and not r.get("signed") else None
        latest_new = max((i for i, c in enumerate(cs) if c.get("kind") == "待签（系统生成）"), default=None)
        keep, drop = [], []
        for i, c in enumerate(cs):
            k = c.get("kind")
            if (k == "续签待签" and i != cur) or (k == "待签（系统生成）" and i != latest_new):
                drop.append(c)
            else:
                keep.append(c)
        changed = bool(drop)
        if cur is not None and cur < len(cs):
            r["contract"] = keep.index(cs[cur])
        a["contracts"] = keep
        for c in keep:
            if c.get("tpl") == TEMPLATE_REV or c.get("kind") not in ("续签待签", "待签（系统生成）"):
                continue
            if c.get("kind") == "续签待签":
                if not r or keep.index(c) != r.get("contract"):
                    continue
                level, start, end, renew = r["level"], dt.date.fromisoformat(r["start"]), dt.date.fromisoformat(r["end"]), True
                r["tpl"] = TEMPLATE_REV
                r.pop("sent", None)
            else:
                if not c.get("expiry"):
                    continue
                end = dt.date.fromisoformat(c["expiry"])
                nxt = end + dt.timedelta(days=1)  # 旧草稿没存开始日：按默认 3 年合同期倒推
                start = dt.date.fromisoformat(c["start"]) if c.get("start") else nxt.replace(year=nxt.year - 3)
                level, renew = a.get("level") if a.get("level") in LEVELS else "standard", False
            fname, data = contracts.generate(a, level, start, end, renew=renew)
            drop.append(dict(c))
            c.update(fid=vault.put_file(data), name=fname, tpl=TEMPLATE_REV)
            c.pop("edited", None)
            changed = True
        if not changed:
            continue
        vault.put("agent", a)
        for c in drop:
            trash_contract(a, c, "旧模板草稿，已换成新模板")
        vault.audit(f"old-template contracts cleaned ({len(drop)})")


def auto_renewals():
    """到期前 90 天自动把续签合同准备好，等 POLO 检查后一键发出。"""
    refresh_old_drafts()
    for a in renewal_candidates():
        if days_left(a["contract_expiry"]) <= AUTO_DRAFT_DAYS and not a.get("renewal"):
            stats = agent_ay_stats(a["id"], last_ay())
            start, end = renewal_period(a)
            draft_renewal(a, suggest_level(a, stats), start, end)
            a["renewal"]["auto"] = True
            vault.put("agent", a)
            vault.audit("renewal auto-drafted")


@app.route("/renewals")
@login_required
def renewals():
    refresh_old_drafts()
    ay, cay = last_ay(), current_ay()
    comms, agents = vault.all("commission"), agents_by_id()
    rows = []
    for a in renewal_candidates():
        stats = agent_ay_stats(a["id"], ay, comms, agents)
        cur = agent_ay_stats(a["id"], cay, comms, agents)
        start, end = renewal_period(a)
        # 等级按上一个完整学年；本学年第一年学生已经满 6 人的也建议高级
        basis = cur if cur["ft1"] >= 6 else stats
        rows.append({"a": a, "stats": stats, "cur": cur, "level": suggest_level(a, basis), "start": start, "end": end,
                     "r": a.get("renewal") or {}})
    counts = {}
    for x in rows:
        lv = x["r"].get("level") or x["level"]
        counts[lv] = counts.get(lv, 0) + 1
    return render_template("renewals.html", rows=rows, ay=ay, cay=cay, counts=counts, auto_days=AUTO_DRAFT_DAYS, renew_days=RENEW_DAYS)


@app.post("/agents/<aid>/renew")
@login_required
def agent_renew(aid):
    a = vault.get("agent", aid) or abort(404)
    level = request.form.get("level") or a.get("level") or "standard"
    r = a.get("renewal") or {}
    if request.form.get("regenerate") or not r or r.get("level") != level:
        start = dt.date.fromisoformat(request.form.get("start") or renewal_period(a)[0].isoformat())
        end = dt.date.fromisoformat(request.form.get("end") or renewal_period(a)[1].isoformat())
        fname, data = draft_renewal(a, level, start, end)
    else:
        c = a["contracts"][r["contract"]]
        fname, data = c["name"], vault.get_file(c["fid"])
    vault.put("agent", a)
    if request.form.get("send"):  # 先到邮件页面预览、翻译、修改，再发
        return redirect(url_for("renew_mail", aid=aid))
    if request.form.get("edit"):
        return redirect(url_for("contract_edit", aid=aid, i=a["renewal"]["contract"], back=url_for("renewals")))
    flash("续签合同已生成，可在线预览检查")
    return redirect(safe_back(request.form.get("back")) or url_for("renewals"))


# 代理邮件可选的翻译语言：开发邮件的 14 种 + 老代理常见的几种
AGENT_LANGS = {k: v["label"] for k, v in i18n.TRANSLATIONS.items()}
AGENT_LANGS.update({"zh": "简体中文", "zh-TW": "繁體中文", "es": AGENT_LANGS.get("es", "Spanish"), "fr": "French",
                    "it": "Italian", "ru": "Russian", "ar": "Arabic", "id": "Indonesian", "mn": "Mongolian",
                    "kk": "Kazakh", "uk": "Ukrainian"})
EXTRA_COUNTRY_LANG = {"china": "zh", "中国": "zh", "taiwan": "zh-TW", "台湾": "zh-TW", "hong kong": "zh-TW", "香港": "zh-TW",
                      "macau": "zh-TW", "spain": "es", "argentina": "es", "ecuador": "es", "venezuela": "es",
                      "france": "fr", "italy": "it", "russia": "ru", "kazakhstan": "kk", "ukraine": "uk",
                      "mongolia": "mn", "indonesia": "id", "saudi arabia": "ar", "uae": "ar", "egypt": "ar",
                      "switzerland": "de", "korea": "ko", "韩国": "ko", "日本": "ja", "越南": "vi", "泰国": "th",
                      "墨西哥": "es", "巴西": "pt", "土耳其": "tr", "俄罗斯": "ru"}


def agent_lang(a):
    """代理用什么语言：先看上次选过的，再按国家猜；英语国家或猜不到就返回空。"""
    if a.get("mail_lang") in AGENT_LANGS:
        return a["mail_lang"]
    country = (a.get("country") or "").strip()
    return i18n.lang_for(country) or EXTRA_COUNTRY_LANG.get(country.lower(), "")


def safe_back(url):
    return url if url and url.startswith("/") and not url.startswith("//") else ""


@app.route("/agents/<aid>/renew/mail", methods=["GET", "POST"])
@login_required
def renew_mail(aid):
    a = vault.get("agent", aid) or abort(404)
    r = a.get("renewal") or {}
    if not r:
        flash("先生成续签合同")
        return redirect(url_for("renewals"))
    c = a["contracts"][r["contract"]]
    if request.method == "POST":
        to = (request.form.get("to") or "").strip()
        if not to:
            flash("请填收件人邮箱")
            return redirect(url_for("renew_mail", aid=aid))
        if request.form.get("lang") in AGENT_LANGS or request.form.get("lang") == "":
            a["mail_lang"] = request.form.get("lang")
        with mailer.Outbox(vault) as box:
            box.send(to, request.form["subject"], request.form["body"], [(c["name"], vault.get_file(c["fid"]))],
                     agent_id=aid, kind="renewal", sender=request.form.get("sender"))
        a["renewal"]["sent"] = dt.date.today().isoformat()
        vault.put("agent", a)
        flash(f"续签合同和续签信已发给 {a['name']}，签回后在代理页面上传即可")
        return redirect(url_for("renewals"))
    tpl = et.fill(et.RENEW_LETTER, contact=a.get("contact") or a["name"], start=r["start"], end=r["end"],
                  kit_link=kit_link())
    return render_template("renew_mail.html", a=a, r=r, c=c, tpl=tpl, langs=AGENT_LANGS, lang=agent_lang(a))


@app.post("/mail/translate")
@login_required
def mail_translate():
    cfg = settings()
    if not ai.available(cfg):
        return {"error": "请先在设置里设置 AI（本地模型）"}, 400
    lang = request.form.get("lang", "")
    if lang not in AGENT_LANGS:
        return {"error": "不支持的语言"}, 400
    try:
        return ai.translate(cfg, request.form["subject"], request.form["body"], AGENT_LANGS[lang])
    except Exception as e:
        return {"error": str(e)[:300]}, 500


# ---------- 代理业绩 ----------
@app.route("/performance")
@login_required
def performance():
    comms, agents = vault.all("commission"), agents_by_id()
    ays = sorted({academic_year(c["term"]) for c in comms} | {current_ay()}, reverse=True)
    ay = request.args.get("ay") or current_ay()
    prev = f"{int(ay[:2]) - 1:02d}-{ay[:2]}" if re.fullmatch(r"\d\d-\d\d", ay) else ""
    studs = vault.all("student")
    today = dt.date.today()
    rows = []
    for a in agents.values():
        if a.get("status") in ("潜在", "洽谈中"):
            continue
        s = agent_ay_stats(a["id"], ay, comms, agents)
        p = agent_ay_stats(a["id"], prev, comms, agents) if prev else {"commission": 0}
        apps = [x for x in studs if x.get("agent_id") == a["id"] and academic_year(x.get("intake") or "") == ay]
        enrolled = sum(1 for x in apps if x.get("stage") == "已入学")
        last = max((c.get("paid_on") or c.get("sent_on") or "" for c in comms if c["agent_id"] == a["id"]), default="")
        last_term = max((c for c in comms if c["agent_id"] == a["id"]), key=lambda c: c.get("term_order", -1), default=None)
        rows.append({"a": a, **s, "prev": p["commission"], "apps": len(apps), "enrolled": enrolled,
                     "last_term": last_term["term"] if last_term else "",
                     "idle": not last or (today - dt.date.fromisoformat(last[:10])).days > 365})
    rows.sort(key=lambda r: (-r["commission"], -r["students"]))
    active = [r for r in rows if r["students"] or r["commission"]]
    return render_template("performance.html", rows=rows, ay=ay, ays=ays, prev=prev, active=active,
                           idle=[r for r in rows if r["idle"] and r["a"].get("status") == "活跃"],
                           total=sum(r["commission"] for r in rows), fees=sum(r["fees"] for r in rows))


# ---------- 宣传资料库 ----------
ASSET_KINDS = ["资料包", "宣传册", "传单", "PPT", "申请表", "文字介绍", "校历", "视频", "照片", "Logo", "其它"]
PARTNER_KIT = ("https://johnathanacademy-my.sharepoint.com/:f:/g/personal/admissionsja_cseec_ca/"
               "EtK0pf0Saj9Dv_j83LS5ELIB7eB9mA0ur8CqfItIUf3c_A")
# 来自学校 OneDrive「宣传」文件夹和招生处共享的 Business partner 文件夹；链接请在 OneDrive 里「共享 → 复制链接」后贴进来
DEFAULT_ASSETS = [
    ("合作伙伴资料包（Business partner 共享文件夹）", "资料包", "中英", PARTNER_KIT,
     "招生处 OneDrive：Important documents for Business partner（申请表、单科申请表、中英文 Flyer、学校介绍 PPT）。Start Letter 和续签信里的资料链接就是它", True),
    ("JA Flyer（英文）", "传单", "英文", "", "宣传 / 宣传册 / JA Flyer E Share.pdf", True),
    ("JA Flyer（中文）", "传单", "中文", "", "宣传 / 宣传册 / JA Flyer 中 Share.pdf", True),
    ("学校画册 2023 英文版（印刷版）", "宣传册", "英文", "", "宣传 / 宣传册 / 画册英文版2023-印刷版.pdf", True),
    ("乔纳森学院画册 中文版", "宣传册", "中文", "", "宣传 / 宣传册 / 乔纳森学院(中文版) Last edit.pdf", True),
    ("学校介绍 PPT 英文 1.3", "PPT", "英文", "", "宣传 / PPT / JA 宣传 Eng1.3 Mar-14-24.pdf", True),
    ("学校介绍 PPT 中文 1.3", "PPT", "中文", "", "宣传 / PPT / JA 宣传 1.3 Mar-14-24.pdf", True),
    ("Admission Application Form", "申请表", "英文", "", "Business partner 文件夹", True),
    ("Single Course Application Form", "申请表", "英文", "", "Business partner 文件夹", True),
    ("学校介绍文字 JA PUBLICITY", "文字介绍", "英文", "", "Agent / Template Docs / JA PUBLICITY.docx（2023 版，数字需要更新）", False),
    ("招生时家长常问的问题汇总", "文字介绍", "中文", "", "宣传 / 在招生时遇到家长问到的问题汇总.docx", False),
    ("学校日历", "校历", "英文", "", "招生处 OneDrive 里是 2024-2025 版，需要换成 2026-27", True),
    ("校长视频 / 宣传视频", "视频", "", "", "宣传 / Principal video、宣传 / Video", False),
    ("校园照片", "照片", "", "", "宣传 / Photoes / Campus、25-26", False),
    ("Logo", "Logo", "", "", "宣传 / LOGO", False),
]


def all_assets():
    rows = vault.all("asset")
    if not rows and not (vault.get("settings", "flags") or {}).get("assets_seeded"):
        for title, kind, lang, url, note, share in DEFAULT_ASSETS:
            vault.put("asset", {"title": title, "kind": kind, "lang": lang, "url": url, "note": note, "share": share,
                                "updated": dt.date.today().isoformat()})
        vault.put("settings", {**(vault.get("settings", "flags") or {"id": "flags"}), "assets_seeded": True})
        rows = vault.all("asset")
    return sorted(rows, key=lambda x: (ASSET_KINDS.index(x["kind"]) if x.get("kind") in ASSET_KINDS else 99, x["title"]))


def kit_link():
    return settings().get("kit_link") or next((x["url"] for x in all_assets() if x.get("kind") == "资料包" and x.get("url")), "")


@app.route("/materials", methods=["GET", "POST"])
@app.route("/materials/<mid>", methods=["GET", "POST"])
@login_required
def materials(mid=None):
    m = vault.get("asset", mid) if mid else None
    if mid and not m:
        abort(404)
    if request.method == "POST":
        if request.form.get("delete") and m:
            vault.delete("asset", mid)
            flash("已删除")
            return redirect(url_for("materials"))
        m = m or {}
        for k in ("title", "kind", "lang", "url", "note"):
            m[k] = request.form.get(k, "").strip()
        m["share"] = bool(request.form.get("share"))
        m["updated"] = dt.date.today().isoformat()
        up = save_upload("file")
        if up:
            m["file"] = up
        vault.put("asset", m)
        flash("已保存")
        return redirect(url_for("materials"))
    rows = all_assets()
    return render_template("materials.html", rows=rows, m=m, kinds=ASSET_KINDS, kit=kit_link(),
                           missing=[x for x in rows if not x.get("url") and not x.get("file")])


# ---------- 年度付款汇总 / T4A ----------
@app.route("/payments")
@login_required
def payments():
    agents = agents_by_id()
    comms = vault.all("commission")
    rows = [enrich(c, agents, comms) for c in comms]
    years = tax.years(rows)
    year = int(request.args.get("year") or (years[1] if len(years) > 1 and dt.date.today().month <= 3 else years[0]))
    summary = tax.summarize(year, rows, agents)
    return render_template("payments.html", summary=summary, year=year, years=years, deadline=tax.deadline(year),
                           total=sum(s["total"] for s in summary), t4a=[s for s in summary if s["t4a"]],
                           ENTITY_TYPES=tax.ENTITY_TYPES, cfg=settings())


@app.route("/payments/<int:year>.xlsx")
@login_required
def payments_export(year):
    agents = agents_by_id()
    comms = vault.all("commission")
    summary = tax.summarize(year, [enrich(c, agents, comms) for c in comms], agents)
    cfg = settings()
    vault.audit("payments export")
    return send_file(io.BytesIO(tax.workbook(year, summary, {"name": cfg.get("payer_name") or "Johnathan Academy",
                                                             "bn": cfg.get("payer_bn") or ""})),
                     download_name=f"JA Agent Payments {year} + T4A.xlsx", as_attachment=True)


@app.post("/payments/<int:year>/send")
@login_required
def payments_send(year):
    """给代理发年度付款明细（选中的代理，或全部当年有付款的代理）。"""
    agents = agents_by_id()
    comms = vault.all("commission")
    summary = tax.summarize(year, [enrich(c, agents, comms) for c in comms], agents)
    pick = set(request.form.getlist("aid"))
    sent = 0
    with mailer.Outbox(vault) as box:
        for s in summary:
            a = s["agent"]
            if (pick and a["id"] not in pick) or not a.get("email"):
                continue
            line = (f"A T4A slip (Box 020, self-employed commissions) for {year} will be issued to you by "
                    f"{tax.deadline(year):%B %d, %Y}." if s["t4a"] else "")
            tpl = et.fill(et.ANNUAL_STATEMENT, contact=a.get("contact") or a["name"], agent=a["name"], year=year,
                          amount=money(s["total"]), t4a_line=line)
            box.send(a["email"], tpl["subject"], tpl["body"],
                     [(f"JA {year} Commission Payments {a.get('code', '')}.xlsx", tax.agent_statement(year, s))],
                     agent_id=a["id"], kind="statement")
            sent += 1
    flash(f"已给 {sent} 个代理发出 {year} 年度付款明细")
    return redirect(url_for("payments", year=year))


# ---------- 设置与安全 ----------
@app.post("/settings/ai-test")
@login_required
def settings_ai_test():
    cfg = dict(settings())
    for k in ("ai_mode", "local_llm_url", "local_model", "local_api_key", "searxng_url", "openai_key", "openai_model", "kit_link", "payer_name", "payer_bn"):  # 用页面上刚填的值测试，不必先保存
        if request.form.get(k):
            cfg[k] = request.form[k].strip()
    try:
        if ai.mode(cfg) == "claude":
            ai.claude_translate(cfg.get("anthropic_key", ""), "Hello", "Hello {name}", "Spanish")
            flash("Claude API 连接正常")
        else:
            models, hits = local_ai.check(cfg)
            flash(f"本地模型连接正常（{models}），搜索返回 {hits} 条结果")
    except Exception as e:
        flash(f"连接失败：{str(e)[:200]}")
    return redirect(url_for("settings_page"))



@app.route("/settings/import", methods=["POST"])
@login_required
def settings_import():
    """首次导入：佣金总表 Excel（代理 + 历史佣金）+ 已签合同文件（加密保存，按文件名里的代理编号和到期年月归档）。"""
    if vault.all("agent") and request.form.get("confirm") != "1":
        flash("系统里已经有代理了，再导入会重复。确实要导入请勾选「我确认要再次导入」")
        return redirect(url_for("settings_page"))
    wb = request.files.get("workbook")
    if not wb or not wb.filename.lower().endswith(".xlsx"):
        flash("请选择佣金总表 Excel（.xlsx）")
        return redirect(url_for("settings_page"))
    fids, skipped = {}, 0
    for f in request.files.getlist("contracts"):
        if f and f.filename:
            name = os.path.basename(f.filename)
            if importer.parse_contract_name(name)[0]:
                fids[name] = vault.put_file(f.read())
            else:
                skipped += 1
    with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
        wb.save(tmp.name)
        try:
            n_agents, n_rows = importer.import_into(vault, tmp.name, contract_fids=fids)
        except Exception as e:
            flash(f"导入失败：{e}")
            return redirect(url_for("settings_page"))
    tracker = request.files.get("tracker")
    n_prospects = import_tracker(tracker.read()) if tracker and tracker.filename else 0
    flash(f"已导入 {n_agents} 个代理、{n_rows} 条历史佣金；{len(fids)} 份合同已归档"
          + (f"；{n_prospects} 个潜在代理" if n_prospects else "")
          + (f"，{skipped} 个文件名里没有代理编号（如 J202103），没有导入，可以到代理页面单独上传" if skipped else ""))
    return redirect(url_for("agents_list"))


def _contract_files(z):
    """包里的合同：[(文件名, 代理编号或 None, 到期日, zip 内路径)]"""
    out = []
    for n in z.namelist():
        base = os.path.basename(n)
        if n.startswith("contracts/") and base:
            code, exp = importer.parse_contract_name(base)
            out.append((base, code, exp, n))
    return out


def _name_key(n):
    """同一个学生不同写法算一个：'Li, jiaheng' = 'Jiaheng Li'；括号里的英文名不算。"""
    n = re.sub(r"\(.*?\)", " ", n or "").replace(",", " ").lower()
    return " ".join(sorted(w for w in re.split(r"\s+", n) if w))


def _student_match(name, pool):
    k = _name_key(name)
    hit = next((s for s in pool if _name_key(s.get("name")) == k), None)
    if hit:
        return hit
    ws = set(k.split())  # 系统里少写了英文名（Mai Bo / Mai Bo Jason）也算同一个人
    return next((s for s in pool if len(ws & set(_name_key(s.get("name")).split())) >= 2
                 and (ws <= set(_name_key(s.get("name")).split()) or set(_name_key(s.get("name")).split()) <= ws)), None)


def _students_pack(z):
    pack = json.loads(z.read("students.json").decode("utf-8"))
    codes = {a.get("code"): a for a in vault.all("agent") if a.get("code")}
    have = vault.all("student")
    comms = vault.all("commission")
    rows = []
    for x in pack.get("students", []):
        a = codes.get(x.get("agent_code"))
        old = _student_match(x["name"], have)
        rows.append(dict(x, agent=a, exists=old, agent_missing=bool(x.get("agent_code") or x.get("agent_hint")) and not a))
    stmts = []
    for st in pack.get("statements", []):
        a = codes.get(st.get("agent_code"))
        cur = next((c for c in comms if a and c["agent_id"] == a["id"] and c["term"] == st["term"]), None)
        r = (pack.get("rosters") or {}).get(st.get("agent_code")) or {}
        stmts.append(dict(st, agent=a, existing=cur, locked=bool(cur and cur.get("locked")),
                          roster=r.get("name"), n=len(st.get("students", []))))
    return pack, rows, stmts


def _students_import(z):
    """导入学生（按名字合并，不重复）、归到代理、建本学期结算单草稿和邮件草稿、存各代理的 Roster 底稿。都不会发信。"""
    pack, rows, stmts = _students_pack(z)
    today = dt.date.today().isoformat()
    by_name = {}
    n_new = n_upd = 0
    for x in rows:
        st = x["exists"] or {"history": {}, "docs": []}
        if not x["exists"]:
            n_new += 1
        else:
            n_upd += 1
        for k in ("name", "grade", "program", "origin", "intake", "courses", "period", "folder"):
            if x.get(k) not in (None, ""):
                st[k] = x[k] if k != "name" or not x["exists"] else st.get("name") or x[k]
        if x.get("tuition") is not None:
            st["tuition"] = x["tuition"]
        if x.get("note"):
            st["notes"] = x["note"]
        st["roster_note"] = x.get("roster_note") or ""
        if not st.get("agent_manual"):  # 你手动分配过的不改
            if x["agent"]:
                st.update(agent_id=x["agent"]["id"], uses_agent=True)
            elif x.get("self"):
                st.update(agent_id="", uses_agent=False)
            else:
                st.update(agent_id="", uses_agent=True, agent_hint=x.get("agent_hint") or x.get("agent_code") or "")
        if x.get("stage") in STUDENT_STAGES and st.get("stage") != x["stage"]:
            st["stage"] = x["stage"]
            st.setdefault("history", {})[x["stage"]] = today
        st["imported_from"] = pack.get("source", "")
        st = vault.put("student", st)
        by_name[x["name"]] = st
    rosters = pack.get("rosters") or {}
    n_stmt, skipped = 0, []
    orders = {t: i for i, t in enumerate(terms())}
    for x in stmts:
        a = x["agent"]
        if not a:
            skipped.append(f"{x.get('agent_hint') or x.get('agent_code')}（系统里没有这个代理）")
            continue
        r = rosters.get(x["agent_code"])
        if r and r.get("file") in z.namelist():
            a = vault.get("agent", a["id"])
            a["roster"] = {"name": r.get("name") or os.path.basename(r["file"]), "fid": vault.put_file(z.read(r["file"])),
                           "updated": today, "agent_label": x.get("agent_label", "")}
            vault.put("agent", a)
        elif x.get("agent_label"):
            a = vault.get("agent", a["id"])
            a["roster"] = dict(a.get("roster") or {}, agent_label=x["agent_label"])
            vault.put("agent", a)
        if x["locked"]:
            skipped.append(f"{a['name']}（{x['term']} 结算单已经发出，没有改）")
            continue
        c = x["existing"] or {"agent_id": a["id"], "term": x["term"], "status": "草稿", "students": [],
                              "term_order": orders.get(x["term"], len(orders) + 100)}
        if x["term"] not in orders:
            vault.put("term", {"name": x["term"], "order": max([0] + [t.get("order", 0) for t in vault.all("term")]
                                                                + [k.get("term_order", 0) for k in vault.all("commission")]) + 1,
                               "created": today})
            orders = {t: i for i, t in enumerate(terms())}
        keep = [s2 for s2 in c.get("students", []) if all(_name_key(s2.get("name")) != _name_key(n) for n in x["students"])]
        add = []
        for n in x["students"]:
            st = by_name.get(n) or {}
            add.append({"name": st.get("name") or n, "program": st.get("program") or "FT1", "courses": st.get("courses") or "",
                        "origin": st.get("origin") or "International", "tuition": st.get("tuition") or 0,
                        "note": "", "period": st.get("roster_note") or st.get("period") or "", "rate": None,
                        "grade": st.get("grade") or "", "student_id": st.get("id")})
        c["students"] = keep + add
        c["imported"] = False
        c["draft_mail"] = {k: v for k, v in {"greet": x.get("greet"), "to": x.get("to")}.items() if v}
        c = vault.put("commission", c)
        for n in x["students"]:
            if by_name.get(n):
                by_name[n]["commission_id"] = c["id"]
                vault.put("student", by_name[n])
        n_stmt += 1
    return n_new, n_upd, n_stmt, skipped


@app.route("/inbox")
@login_required
def inbox_page():
    """加密收件箱：先预览包里的内容（特别是收款信息），确认后才导入。"""
    bs = inbox.bundles(vault)
    sel = request.args.get("b") or next((b["name"] for b in bs if b["complete"]), None)
    pv, err = None, None
    if sel and any(b["name"] == sel and b["complete"] for b in bs):
        try:
            with inbox.open_bundle(vault, sel) as z:
                names = z.namelist()
                agents, records = [], []
                if "workbook.xlsx" in names:
                    with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
                        tmp.write(z.read("workbook.xlsx"))
                        tmp.flush()
                        ag, records = importer.read_workbook(tmp.name)
                    have = {a.get("code"): a for a in vault.all("agent") if a.get("code")}
                    for a in ag.values():
                        old = have.get(a.get("code"))
                        a["exists"] = bool(old)
                        a["pay_changed"] = bool(old) and any(
                            (old.get(k) or "") != (a.get(k) or "") for k in ("etransfer", "payee", "detail"))
                        agents.append(a)
                    agents.sort(key=lambda a: a.get("code") or "~" + a["name"])
                tracker_n = 0
                if "tracker.xlsx" in names:
                    wb = openpyxl.load_workbook(io.BytesIO(z.read("tracker.xlsx")), read_only=True)
                    tracker_n = sum(max(ws.max_row - 1, 0) for ws in wb.worksheets)
                cf = _contract_files(z)
                if "students.json" in names:
                    spack, srows, sstmts = _students_pack(z)
                    spv = {"pack": spack, "rows": srows, "stmts": sstmts,
                           "agents": agents_by_id()}
                else:
                    spv = None
                pv = {"agents": agents, "rows": len(records), "contracts": cf, "students": spv,
                      "matched": sum(1 for c in cf if c[1]), "tracker_n": tracker_n,
                      "terms": sorted({r["term"] for r in records}, key=lambda t: next(
                          (r["term_order"] for r in records if r["term"] == t), 0))}
        except Exception as e:
            err = str(e)
    return render_template("inbox.html", bundles=bs, sel=sel, pv=pv, err=err)


@app.post("/inbox/<name>/import")
@login_required
def inbox_import(name):
    with inbox.open_bundle(vault, name) as z:
        has_workbook = "workbook.xlsx" in z.namelist()
    if has_workbook and vault.all("agent") and request.form.get("confirm") != "1":
        flash("系统里已经有代理了，再导入会重复。确实要导入请勾选「我确认要再次导入」")
        return redirect(url_for("inbox_page", b=name))
    try:
        with inbox.open_bundle(vault, name) as z:
            fids = {base: vault.put_file(z.read(path)) for base, code, _, path in _contract_files(z) if code}
            msg = []
            if "workbook.xlsx" in z.namelist():
                with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
                    tmp.write(z.read("workbook.xlsx"))
                    tmp.flush()
                    n_agents, n_rows = importer.import_into(vault, tmp.name, contract_fids=fids)
                msg.append(f"{n_agents} 个代理、{n_rows} 条历史佣金、{len(fids)} 份已签合同")
            if "tracker.xlsx" in z.namelist():
                msg.append(f"{import_tracker(z.read('tracker.xlsx'))} 个潜在代理")
            if "students.json" in z.namelist():
                n_new, n_upd, n_stmt, skipped = _students_import(z)
                msg.append(f"学生新增 {n_new} 名、更新 {n_upd} 名，{n_stmt} 张结算单草稿（含邮件和附件草稿，都没有发出）"
                           + ("；没建结算单：" + "、".join(skipped) if skipped else ""))
    except Exception as e:
        flash(f"导入失败：{e}")
        return redirect(url_for("inbox_page", b=name))
    inbox.remove(vault, name)
    vault.audit("inbox bundle imported")
    flash("已导入：" + "，".join(msg) + "。收件箱里的加密包已删除。请检查，不对的可以直接改。")
    return redirect(url_for("agents_list") if has_workbook else url_for("students"))


@app.post("/inbox/<name>/delete")
@login_required
def inbox_delete(name):
    try:
        inbox.remove(vault, name)
    except (ValueError, OSError):
        abort(404)
    vault.audit("inbox bundle discarded")
    flash("已删除这个包，没有导入任何内容")
    return redirect(url_for("inbox_page"))


@app.route("/settings", methods=["GET", "POST"])
@login_required
def settings_page():
    cfg = settings()
    if request.method == "POST":
        f = request.form
        if f.get("new_password"):
            if len(f["new_password"]) < 12 or not vault.try_unwrap(f.get("old_password", "")):
                flash("原密码错误或新密码少于 12 位")
            else:
                vault.change_password(f["new_password"])
                _trusted_save({})  # 换密码时，所有记住的设备都要重新输验证码
                flash("主密码已更换；所有记住的设备下次登录要重新输验证码")
        else:
            if f.get("m365_form"):
                cfg["digest_on"] = "1" if f.get("digest_on") else "0"
                if (f.get("m365_client_id", "").strip() != cfg.get("m365_client_id", "")
                        or f.get("m365_tenant", "").strip() != cfg.get("m365_tenant", "")):
                    cfg.pop("m365_accounts", None)  # 换了应用或租户，所有邮箱都要重新授权
                    cfg.pop("mail_route", None)
            if f.get("route_form"):
                users = {x["user"] for x in mailer.accounts(cfg)}
                cfg["mail_route"] = {k: f.get("route_" + k) for k, _ in mailer.MAIL_KINDS if f.get("route_" + k) in users}
            for k in ("app_senders", "m365_tenant", "m365_client_id", "digest_to",
                      "smtp_host", "smtp_port", "smtp_user", "from_addr", "imap_host", "imap_user", "accountant_email",
                      "ai_mode", "local_llm_url", "local_model", "searxng_url", "openai_model"):
                if k in f:
                    cfg[k] = f.get(k, "").strip()
            for k in ("smtp_password", "imap_password", "anthropic_key", "local_api_key", "openai_key"):
                if f.get(k):
                    cfg[k] = f[k].strip()
            vault.put("settings", cfg)
            flash("设置已加密保存")
        return redirect(url_for("settings_page"))
    log = [(dt.datetime.fromtimestamp(t).strftime("%Y-%m-%d %H:%M"), e) for t, e in vault.audit_log()]
    return render_template("settings.html", cfg=cfg, app_on=mailer.app_on(), log=log, trusted_n=len(_trusted_load()), this_trusted=device_trusted(), portal_on=portal.enabled(), bk=backup.status(DATA_DIR), bk_dir=backup.backup_dir(),
                           bk_files=backup.listing()[:5])


@app.post("/settings/m365/start")
@login_required
def m365_start():
    cfg = settings()
    if not cfg.get("m365_client_id"):
        return {"error": "请先填应用程序 ID 并保存"}
    try:
        r = mailer.m365_start(cfg)
    except Exception as e:  # noqa: BLE001
        return {"error": f"连不上微软登录服务：{type(e).__name__}"}
    return {k: r.get(k) for k in ("user_code", "verification_uri", "device_code", "interval")}


@app.post("/settings/m365/poll")
@login_required
def m365_poll_route():
    try:
        state = mailer.m365_poll(vault, request.form.get("device_code", ""))
    except Exception as e:  # noqa: BLE001
        state = f"{type(e).__name__}"
    if state == "ok":
        vault.audit("Microsoft 365 mail connected")
    return {"state": state}


@app.post("/settings/m365/disconnect")
@login_required
def m365_disconnect():
    user = request.form.get("user", "")
    mailer.m365_remove(vault, user)
    vault.audit("Microsoft 365 mailbox disconnected")
    left = mailer.accounts(settings())
    flash(f"已断开 {user}" + ("" if left else "，回到演示模式") + ("；原来分给它的邮件改用默认邮箱发" if left else ""))
    return redirect(url_for("settings_page"))


@app.post("/settings/m365/test")
@login_required
def m365_test():
    user = request.form.get("user") or settings().get("m365_user")
    try:
        with mailer.Outbox(vault) as box:
            box.send(user, "代理管理系统 · 测试邮件", f"这是一封测试邮件，说明系统已经可以用 {user} 发信了。", sender=user)
        flash(f"测试邮件已从 {user} 发给它自己，请查收")
    except Exception as e:  # noqa: BLE001
        flash(f"发送失败：{e}")
    return redirect(url_for("settings_page"))


@app.post("/settings/trusted/clear")
@login_required
def trusted_clear():
    _trusted_save({})
    vault.audit("trusted devices cleared")
    flash("已取消所有记住的设备，下次登录都要输验证码")
    resp = redirect(url_for("settings_page"))
    resp.delete_cookie(TRUST_COOKIE)
    return resp


@app.post("/settings/backup/verify")
@login_required
def settings_backup_verify():
    try:
        r = backup.verify(vault)
        msg = f"恢复演练：{r['file']} 里 {r['records']} 条记录、{r['files']} 个文件都能用现在的密钥解开"
        flash(msg + ("。" if not r["bad"] else f"，但有 {r['bad']} 项解不开，请告诉 Claude"))
        vault.audit(f"backup verify: {r['records']} records, {r['files']} files, {r['bad']} bad")
    except Exception as e:  # noqa: BLE001
        flash(f"恢复演练失败：{e}")
    return redirect(url_for("settings_page"))


@app.post("/settings/backup")
@login_required
def settings_backup():
    try:
        path = backup.run(DATA_DIR)
        flash(f"备份完成：{os.path.basename(path)}（全部是密文）")
    except Exception as e:
        flash(f"备份失败：{e}")
    return redirect(url_for("settings_page"))


def weekly_digest(force=False):
    """每周一早上给 POLO 发一封汇总：快到期的合同、没走完的结算单、要跟进的潜在代理、新收到的代理邮件。"""
    cfg = settings()
    today = dt.date.today()
    if not force and (today.weekday() != 0 or dt.datetime.now().hour < 8 or cfg.get("digest_on", "1") == "0"
                      or cfg.get("digest_last") == today.isoformat()):
        return False
    to = cfg.get("digest_to") or cfg.get("m365_user") or cfg.get("smtp_user")
    if not mailer.mail_on(cfg) or not to:
        return False
    agents = agents_by_id()
    comms = vault.all("commission")
    rows = [enrich(c, agents, comms) for c in comms if not c.get("imported")]
    lines = [f"代理管理系统 · 每周汇总（{today.isoformat()}）", ""]
    exp = sorted([a for a in agents.values() if a.get("contract_expiry") and a.get("status") in ("活跃", "已签约")
                  and days_left(a["contract_expiry"]) <= 90], key=lambda a: a["contract_expiry"])
    lines.append(f"一、90 天内到期的合同：{len(exp)} 个")
    for a in exp[:15]:
        r = a.get("renewal") or {}
        st = "已签回" if r.get("signed") else "续签已发出" if r.get("sent") else "续签草稿待检查" if r else "还没准备续签"
        lines.append(f"   • {a.get('name')} {a.get('code') or ''} — {a['contract_expiry']}（{st}）")
    lines += ["", "二、还没走完的结算单"]
    for stage, label in (("已发结算单", "已发出、等代理确认"), ("代理已确认", "代理已确认、还没交给会计"),
                         ("已交会计", "已交会计、还没付款")):
        xs = [r for r in rows if r["stage"] == stage]
        if xs:
            lines.append(f"  {label}：{len(xs)} 张，共 ${sum(r['amount'] for r in xs):,.2f}")
            for r in xs[:10]:
                lines.append(f"   • {r.get('agent_name') or agents.get(r['agent_id'], {}).get('name', '')} · {r['term']} · ${r['amount']:,.2f}")
    fu = [p for p in vault.all("prospect") if needs_followup(p)]
    lines += ["", f"三、需要跟进的潜在代理：{len(fu)} 个"]
    week_ago = (dt.datetime.now() - dt.timedelta(days=7)).isoformat()
    inbox_n = sum(1 for m in vault.all("mail") if m["direction"] == "in" and m.get("at", "") >= week_ago)
    lines += [f"四、过去 7 天收到的代理邮件：{inbox_n} 封", "", "详情请在学校网络里打开 https://192.168.1.75:8443 查看。"]
    with mailer.Outbox(vault) as box:
        box.send(to, f"代理管理 · 每周汇总 {today.isoformat()}", "\n".join(lines), kind="digest")
    rec = vault.get("settings", "mail") or {"id": "mail"}
    rec["digest_last"] = today.isoformat()
    vault.put("settings", rec)
    vault.audit("weekly digest sent")
    return True


def start_jobs():
    mailer.start_scheduler(vault)
    backup.start(DATA_DIR)

    def daily():
        while True:
            try:
                if vault.unlocked:
                    auto_renewals()
                    purge_trash()
                    weekly_digest()
                    sync_admissions()
            except Exception as e:
                vault.audit(f"renewal job error: {type(e).__name__}")
            time.sleep(3600)
    threading.Thread(target=daily, daemon=True).start()


# ---------- 代理自助查询页（默认关闭）----------
portal_view = portal.register(app, vault, {"agents_by_id": agents_by_id, "enrich": enrich, "all_assets": all_assets,
                                           "statement_file": _statement_file, "save_upload": save_upload,
                                           "mailer": mailer, "et": et})


@app.route("/agents/<aid>/portal-preview")
@login_required
def portal_preview(aid):
    """POLO 预览某个代理在自助页上看到的内容（自助页未开放时也能看）。"""
    return portal_view(vault.get("agent", aid) or abort(404), preview=True)


if __name__ == "__main__":
    start_jobs()
    # 只监听本机；对外请用 Nginx/Caddy + HTTPS 反向代理
    app.run(host=os.environ.get("HOST", "127.0.0.1"), port=int(os.environ.get("PORT", 8800)))
