"""一份合同模板，按代理等级自动生成 Service Contract。

原来的两份合同（20,10,5% 和 25,10,5%）正文只差第一年比例这两行，
所以只保留一份母版（contract_templates/service_contract_master.docx），
生成时自动填：第一年比例、合作期限、代理编号、Party B 信息、付款信息。
"""
import calendar
import hashlib
import shutil
import subprocess
import tempfile
import threading
import copy
import datetime as dt
import io
import os
import re

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from commission import LEVELS

MASTER = os.path.join(os.path.dirname(__file__), "contract_templates", "service_contract_master.docx")


def _replace(paragraph, pattern, repl):
    """在一个段落里做正则替换，跨 run 也能替换，尽量保留原格式。"""
    runs = paragraph.runs
    text = "".join(r.text for r in runs)
    m = re.search(pattern, text)
    if not m:
        return False
    start, end = m.span()
    new = m.expand(repl) if isinstance(repl, str) else repl(m)
    pos = 0
    placed = False
    for r in runs:
        r_start, r_end = pos, pos + len(r.text)
        pos = r_end
        if r_end <= start or r_start >= end:
            continue
        keep_before = r.text[:max(0, start - r_start)]
        keep_after = r.text[max(0, end - r_start):] if r_end > end else ""
        r.text = keep_before + ("" if placed else new) + keep_after
        placed = True
    return True


def _set_cell(cell, value):
    if value is None:
        return
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = str(value)
        for r in p.runs[1:]:
            r.text = ""
    else:
        # 空格子的段落标记带着标题样式（粗体大字），填进去的字要和甲方那边一样：12 号、不加粗
        r = p.add_run(str(value))
        r.bold = False
        r.font.size = Pt(12)


def _ap_section(d, s, e, rate):
    """AP 课程佣金单列一项「c.」（POLO 2026-11-01 要求）：合同期 + 比例。比例默认 10%，
    写成下划线填空，POLO 可以在线改或打印后手填；rate 为空时只留空格。"""
    ps = d.paragraphs
    ap = next((p for p in ps if p.text.strip().startswith("AP Program:")), None)
    pt = next((p for p in ps if "for a part-time student" in p.text), None)
    if ap is None or pt is None:
        return
    head = copy.deepcopy(pt._p)  # 同一个编号列表，自动排成 c.
    ap._p.addprevious(head)
    hp = Paragraph(head, ap._parent)
    for r in hp.runs[1:]:
        r._r.getparent().remove(r._r)
    hp.runs[0].text = (f"Starting from {s} to {e} if Party B has completed the standard application process "
                       "of Party A for a student taking AP (Advanced Placement) courses.")
    runs = ap.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    runs[0].text = "Party A will pay Party B "
    rpr = runs[0]._r.rPr

    def add(text):  # 新加的字沿用第一段字的字体字号
        r = ap.add_run(text)
        if rpr is not None:
            r._r.insert(0, copy.deepcopy(rpr))
        return r
    blank = add(f"   {rate:g}   " if rate not in (None, "") else "            ")
    blank.underline = True
    blank.bold = True
    add("% of the AP course tuition fee that is paid by the student the first time. "
        "The tuition fee for the following terms is not applicable.")


def _is_blank(el):
    return el.tag == qn("w:p") and not "".join(t.text or "" for t in el.iter(qn("w:t"))).strip() \
        and el.find(".//" + qn("w:drawing")) is None and el.find(".//" + qn("w:pict")) is None \
        and el.find(".//" + qn("w:sectPr")) is None


def _field(p, instr):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def _layout(d):
    """双方签名行放进页脚（每页最下面都有），页脚最后加「Page X of Y」。
    母版里签名行是正文里的几个小表格，靠空行顶到页底，换台电脑或改几个字就会跑到下一页顶上。"""
    sigs = [t for t in d.tables if len(t.rows) == 1 and t.rows[0].cells[0].text.strip().startswith("Signature of Party A")]
    if not sigs:
        return
    keep = copy.deepcopy(sigs[0]._tbl)
    for t in sigs:
        el = t._tbl
        prev = el.getprevious()
        while prev is not None and _is_blank(prev):  # 顶签名行用的空行也去掉
            nxt = prev.getprevious()
            prev.getparent().remove(prev)
            prev = nxt
        el.getparent().remove(el)
    # 乙方义务和附表 A 各从新的一页开始，和原合同一样
    sched = False
    for p in d.paragraphs:
        txt = p.text.strip()
        if txt == "Party B:" or txt.startswith("Schedule A"):
            p.paragraph_format.page_break_before = True
        if txt.startswith("If there is any dispute"):  # 小标题不要单独留在页底
            p.paragraph_format.keep_with_next = True
        sched = sched or txt.startswith("Schedule A")
        if sched and _is_blank(p._p) and p._p.getnext() is not None and p._p.getnext().tag == qn("w:p"):  # 附表 A 小标题前的空行去掉，整页放得下
            p._p.getparent().remove(p._p)
    for s in d.sections:
        ft = s.footer
        ft._element.insert(0, keep)
        gap = OxmlElement("w:p")
        keep.addnext(gap)
        p = ft.add_paragraph()
        p.alignment = 1
        for txt, instr in (("Page ", "PAGE"), (" of ", "NUMPAGES")):
            r = p.add_run(txt)
            r.font.size = Pt(9)
            _field(p, instr)
        for r in p._p.iter(qn("w:r")):
            rpr = r.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                r.insert(0, rpr)
            if rpr.find(qn("w:sz")) is None:
                sz = OxmlElement("w:sz")
                sz.set(qn("w:val"), "18")
                rpr.append(sz)
        break


def default_period(start=None, years=3):
    start = start or dt.date.today()
    end_y, end_m = start.year + years, start.month - 1 or 12
    if start.month == 1:
        end_y -= 1
    end = dt.date(end_y, end_m, calendar.monthrange(end_y, end_m)[1])
    return start, end


def generate(agent, level, start, end, master=MASTER, renew=False):
    """返回 (文件名, docx 字节)。续签用同一份母版，只是合同期和文件名不同。
    AP 课程佣金比例用 agent["ap_rate"]（没设过 = 等级默认 10%，空字符串 = 合同上留空手填）。"""
    rates = LEVELS[level]
    d = docx.Document(master)
    s, e = start.strftime("%B %d, %Y"), end.strftime("%B %d, %Y")
    for p in d.paragraphs:
        _replace(p, r"Starting from 20 _+\s*to 20 _+", f"Starting from {s} to {e}")
        if "for the first academic year" in p.text:
            if "1-5" in p.text:
                _replace(p, r"^\s*\d+\s*%", f"{rates['y1']:g}%")
            elif "6 or more" in p.text:
                _replace(p, r"^\s*\d+\s*%", f"{rates['y1_6']:g}%")
    ap_rate = agent.get("ap_rate", rates.get("ap", 10))
    _ap_section(d, s, e, float(ap_rate) if ap_rate not in (None, "") else None)
    t = d.tables
    name = agent.get("legal_name") or agent.get("name", "")
    _set_cell(t[1].rows[0].cells[1], name)
    # 地址、电话、传真、网站、手机留空，由代理自己填写（POLO 2026-10-27 要求）
    for r, c in ((1, 1), (2, 1), (2, 3), (3, 1), (3, 3)):
        _set_cell(t[1].rows[r].cells[c], "")
    _set_cell(t[1].rows[4].cells[1], agent.get("email", ""))
    _set_cell(t[2].rows[0].cells[1], agent.get("code", ""))
    _set_cell(t[2].rows[1].cells[1], start.isoformat())
    _set_cell(t[2].rows[1].cells[3], end.isoformat())
    _set_cell(t[7].rows[0].cells[1], name)
    _set_cell(t[7].rows[1].cells[1], agent.get("contact", ""))
    # Schedule A：已知的付款信息先填上，其余由代理填写
    if agent.get("payee"):
        _set_cell(t[9].rows[2].cells[2], agent["payee"])
    if agent.get("etransfer"):
        _set_cell(t[9].rows[3].cells[2], agent["etransfer"])
    _layout(d)
    buf = io.BytesIO()
    d.save(buf)
    short = re.sub(r"[\\/:*?\"<>|]", "", agent.get("name", "Agent")).strip()
    fname = f"Service Contract {'Renew ' if renew else ''}({short}) {agent.get('code', '')} until {end.strftime('%Y %b')}.docx"
    return fname, buf.getvalue()


# ---------- 在线预览 / 编辑 ----------
# 按文档顺序把每个段落（含表格里的）编号，预览和保存用同一套编号。
def _walk(parent):
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for el in parent.iter_inner_content():
        if isinstance(el, Paragraph):
            yield ("p", el)
        elif isinstance(el, Table):
            rows = []
            for row in el.rows:
                seen, cells = set(), []
                for cell in row.cells:  # 合并单元格会重复出现，只取一次
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    span = sum(1 for c in row.cells if c._tc is cell._tc)
                    cells.append((span, list(_walk(cell))))
                rows.append(cells)
            yield ("t", rows)


def _paragraphs(blocks):
    for kind, x in blocks:
        if kind == "p":
            yield x
        else:
            for row in x:
                for _, inner in row:
                    yield from _paragraphs(inner)


def _runs_info(p):
    out = []
    for r in p.runs:
        if not r.text:
            continue
        size = r.font.size.pt if r.font.size else None
        st = r.style.font if r.style is not None else None  # 粗体也可能来自字符样式（如 Strong）
        out.append({"t": r.text, "b": bool(r.bold if r.bold is not None else st and st.bold),
                    "i": bool(r.italic if r.italic is not None else st and st.italic),
                    "u": bool(r.underline if r.underline is not None else st and st.underline), "size": size})
    return out


def outline(data):
    """给编辑页用：[{kind:'p', i, text, runs, editable, ...} | {kind:'t', rows:[[{span, blocks}]]}]"""
    d = docx.Document(io.BytesIO(data))
    counter = iter(range(10 ** 6))

    def conv(blocks):
        out = []
        for kind, x in blocks:
            if kind == "p":
                style = (x.style.name if x.style is not None else "") or ""
                indent = x.paragraph_format.left_indent
                out.append({"kind": "p", "i": next(counter), "text": x.text, "runs": _runs_info(x),
                            # 带超链接等特殊内容的段落，文字不全在 run 里，只能看不能改
                            "editable": "".join(r.text for r in x.runs) == x.text,
                            "center": x.alignment is not None and int(x.alignment) == 1,
                            "right": x.alignment is not None and int(x.alignment) == 2,
                            "indent": round(indent.pt) if indent else 0,
                            "head": style.lower().startswith(("heading", "title")),
                            "list": x._p.pPr is not None and x._p.pPr.numPr is not None})
            else:
                out.append({"kind": "t", "rows": [[{"span": s, "blocks": conv(b)} for s, b in row] for row in x]})
        return out
    return conv(list(_walk(d)))


def _edit_paragraph(p, new):
    """只替换改动的那一段文字（前后相同的部分不动），所以粗体、下划线等格式都保留。"""
    runs = p.runs
    old = "".join(r.text for r in runs)
    if old != p.text or new == old:
        return False
    if not runs:
        p.add_run(new)
        return True
    n = min(len(old), len(new))
    a = 0
    while a < n and old[a] == new[a]:
        a += 1
    b = 0
    while b < n - a and old[len(old) - 1 - b] == new[len(new) - 1 - b]:
        b += 1
    start, end, repl = a, len(old) - b, new[a:len(new) - b]
    pos, placed = 0, False
    for r in runs:
        r_start, r_end = pos, pos + len(r.text)
        pos = r_end
        if start == end:  # 纯插入：接在插入点前面那段文字后面，沿用它的格式
            if not placed and (r_start < start <= r_end or start == 0):
                cut = start - r_start
                r.text = r.text[:cut] + repl + r.text[cut:]
                placed = True
            continue
        if r_end <= start or r_start >= end:
            continue
        keep_before = r.text[:max(0, start - r_start)]
        keep_after = r.text[max(0, end - r_start):] if r_end > end else ""
        r.text = keep_before + ("" if placed else repl) + keep_after
        placed = True
    return placed


def apply_edits(data, edits):
    """edits = {段落编号: 新文字}。"""
    d = docx.Document(io.BytesIO(data))
    for i, p in enumerate(_paragraphs(list(_walk(d)))):
        if edits.get(i) is not None:
            _edit_paragraph(p, edits[i])
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------- 排版预览：用 LibreOffice 转成 PDF，再转成图片，看起来和 Word 一样 ----------
_render_lock = threading.Lock()
_pdf_cache = {}


def to_pdf(data):
    """返回 PDF 字节；服务器上没装 LibreOffice 时返回 None。临时文件放在内存盘 /dev/shm，不落到硬盘。"""
    key = hashlib.sha256(data).hexdigest()
    if key in _pdf_cache:
        return _pdf_cache[key]
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    base = "/dev/shm" if os.path.isdir("/dev/shm") else None
    with _render_lock, tempfile.TemporaryDirectory(dir=base) as tmp:
        src = os.path.join(tmp, "c.docx")
        with open(src, "wb") as f:
            f.write(data)
        subprocess.run([soffice, f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--convert-to", "pdf",
                        "--outdir", tmp, src], capture_output=True, timeout=120)
        out = os.path.join(tmp, "c.pdf")
        if not os.path.exists(out):
            return None
        with open(out, "rb") as f:
            pdf = f.read()
    if len(_pdf_cache) > 8:
        _pdf_cache.clear()
    _pdf_cache[key] = pdf
    return pdf


_pdfium_lock = threading.Lock()  # pdfium 不是线程安全的，多页图片同时请求会崩溃，所以一次只渲染一页
_png_cache = {}


def page_count(pdf):
    import pypdfium2
    with _pdfium_lock:
        doc = pypdfium2.PdfDocument(pdf)
        try:
            return len(doc)
        finally:
            doc.close()


def page_png(pdf, n, scale=1.6):
    import pypdfium2
    key = (hashlib.sha256(pdf).hexdigest(), n, scale)
    if key in _png_cache:
        return _png_cache[key]
    with _pdfium_lock:
        doc = pypdfium2.PdfDocument(pdf)
        try:
            page = doc[n]
            img = page.render(scale=scale).to_pil()
            page.close()
        finally:
            doc.close()
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=True)
    if len(_png_cache) > 40:
        _png_cache.clear()
    _png_cache[key] = buf.getvalue()
    return _png_cache[key]
