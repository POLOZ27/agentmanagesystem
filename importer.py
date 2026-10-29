"""从现有《代理表》Excel 和合同文件名导入数据到加密保险库。

用法： python importer.py 代理表.xlsx [合同文件名列表.txt]
（会提示输入主密码，导入后数据只以密文保存）
"""
import calendar
import getpass
import re
import sys

import openpyxl

from commission import parse_scheme

FIELD = {
    "学期缴费全日制学生": "stu_term", "学期已缴费用": "fee_term",
    "全日制学生": "stu_ft", "全日制学生 Year One": "stu_y1", "全日制学生 Year Two": "stu_y2",
    "全日制学生 Year Three": "stu_y3", "全日制学生已缴费用": "fee_ft",
    "兼职学生15%": "stu_pt", "兼职学生已缴费用": "fee_pt",
    "AP 课程": "stu_ap", "AP 返佣比例": "ap_rate_text", "AP 已缴费用": "fee_ap",
    "竞赛报名反10%": "competition", "已缴学费总金额": "excel_total", "其它": "other_note",
    "Commission 金额": "excel_commission", "Commission 学费金额": "excel_commission",
    "E-Transfer address": "etransfer", "支票抬头": "payee", "Detail": "detail", "状态": "status",
    "代理名称": "name", "代理编号": "code", "Commission %": "scheme",
}
MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_abbr) if m}
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")


def _num(v):
    return float(v) if isinstance(v, (int, float)) else None


def parse_contract_name(fn):
    code = re.search(r"J20\d{4}", fn)
    exp = re.search(r"until\s+(\d{4})\s*([A-Za-z]+)?", fn, re.I)
    expiry = None
    if exp:
        y, m = int(exp.group(1)), MONTHS.get((exp.group(2) or "")[:3].lower(), 12)
        expiry = f"{y}-{m:02d}-{calendar.monthrange(y, m)[1]:02d}"
    return (code.group(0) if code else None), expiry


def read_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    agents, records = {}, []
    sheets = wb.worksheets
    for order, ws in enumerate(sheets):
        header = [FIELD.get((c.value or "").strip()) if isinstance(c.value, str) else None for c in ws[1]]
        is_template = ws.title.lower().startswith("form template")
        for row in ws.iter_rows(min_row=2, values_only=True):
            r = {k: v for k, v in zip(header, row) if k and v not in (None, "")}
            name = str(r.get("name") or "").strip()
            if not name:
                continue
            code = str(r.get("code") or "").strip() or None
            key = code or name
            a = agents.setdefault(key, {"code": code, "name": name, "status": "活跃",
                                        "first_term": None if is_template else ws.title})
            a["name"] = name
            if not is_template and not a.get("first_term"):
                a["first_term"] = ws.title
            for k in ("scheme", "etransfer", "payee", "detail"):
                if r.get(k) not in (None, ""):
                    a[k] = r[k]
            if is_template:
                continue
            fees = {k: _num(r.get(k)) for k in ("fee_term", "fee_ft", "fee_pt", "fee_ap", "excel_commission")}
            if not any(fees.values()):
                continue
            rec = {"agent_key": key, "term": ws.title, "term_order": order, "imported": True,
                   "rates": parse_scheme(r.get("scheme")), "scheme_text": str(r.get("scheme") or ""),
                   "status": "已付" if str(r.get("status", "")).strip() == "已付" else "历史导入",
                   "fee_term": fees["fee_term"], "fee_pt": fees["fee_pt"], "fee_ap": fees["fee_ap"],
                   "excel_commission": fees["excel_commission"], "excel_total": _num(r.get("excel_total"))}
            for k in ("stu_term", "stu_ft", "stu_y1", "stu_y2", "stu_y3", "stu_pt", "stu_ap",
                      "ap_rate_text", "competition"):
                if r.get(k) is not None:
                    rec[k] = str(r[k])
            # Excel 只有一栏"全日制已缴费用"：按填了学生的年级归入对应年份
            years = [y for y in ("y2", "y3") if rec.get("stu_" + y) and not rec.get("stu_y1")]
            rec["fee_" + (years[0] if len(years) == 1 else "y1")] = fees["fee_ft"]
            # 早期表格只有总额没有分项：视为全日制 Year One
            if not any(rec.get(k) for k in ("fee_term", "fee_y1", "fee_y2", "fee_y3", "fee_pt", "fee_ap")) \
                    and rec.get("excel_total") and not rec.get("stu_pt"):
                rec["fee_y1"] = rec["excel_total"]
            other = r.get("other_note")
            if isinstance(other, (int, float)):
                rec["other_amount"] = float(other)
            elif other:
                rec["other_note"] = str(other)
            if r.get("ap_rate_text"):
                m = re.findall(r"\d+(?:\.\d+)?", str(r["ap_rate_text"]))
                rec["ap_rate"] = float(m[0]) if len(m) == 1 else None
            records.append(rec)
    for a in agents.values():
        a["rates"] = parse_scheme(a.get("scheme"))
        emails = EMAIL.findall(str(a.get("etransfer", "")))
        a["email"] = emails[0] if emails else ""
        a["scheme_text"] = str(a.pop("scheme", "") or "")
    return agents, records


def import_into(vault, xlsx_path, contract_names=(), contract_fids=None):
    """contract_fids：{文件名: 已加密保存的文件 id}，网页上传合同时用。"""
    agents, records = read_workbook(xlsx_path)
    contracts = {}
    for fn in list(contract_names) + list(contract_fids or {}):
        code, expiry = parse_contract_name(fn)
        if code:
            contracts.setdefault(code, []).append({"name": fn, "expiry": expiry, "fid": (contract_fids or {}).get(fn),
                                                   "kind": "已签"})
    ids = {}
    recent = {r["agent_key"] for r in records if r["term_order"] >= max(x["term_order"] for x in records) - 3}
    today = __import__("datetime").date.today().isoformat()
    for key, a in agents.items():
        cs = sorted(contracts.get(a["code"], []), key=lambda c: c["expiry"] or "")
        a["contracts"] = cs
        a["contract_expiry"] = cs[-1]["expiry"] if cs else None
        a["source"] = "Excel 导入"
        # 合同有效或最近 4 个学期有佣金 => 活跃
        a["status"] = "活跃" if (a["contract_expiry"] or "") >= today or key in recent else "不活跃"
        ids[key] = vault.put("agent", a)["id"]
    for rec in records:
        rec["agent_id"] = ids[rec.pop("agent_key")]
        vault.put("commission", rec)
    vault.audit(f"imported {len(agents)} agents / {len(records)} commission rows")
    return len(agents), len(records)


if __name__ == "__main__":
    from app import vault
    pw = getpass.getpass("主密码: ")
    dek = vault.try_unwrap(pw)
    if not dek:
        sys.exit("密码错误")
    vault.unlock(dek)
    names = open(sys.argv[2], encoding="utf-8").read().splitlines() if len(sys.argv) > 2 else []
    print(import_into(vault, sys.argv[1], [n for n in names if n.strip()]))
